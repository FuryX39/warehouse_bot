"""Отчёт по продажам: заказы выбранного маркетплейса за период.

Сумма = «Цена с НДС» (unit_price) × количество. Если цена в строке не заполнена
(сейчас так на WB), сумма по артикулу остаётся пустой.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.catalog_repository import CatalogRepository
from app.sales_analysis import MARKETPLACE_IDS, MARKETPLACES, marketplace_title
from app.warehouse_order_pricing import parse_money, round_money
from app.warehouse_orders_repository import (
    ORDER_CANCELLED,
    WarehouseOrder,
    WarehouseOrderLine,
    WarehouseOrdersRepository,
)

_TZ = ZoneInfo("Europe/Moscow")
_HEADERS = ("Артикул", "Название", "Количество", "Сумма")
MISSING_NAME = "Товар не найден в каталоге"


@dataclass(frozen=True)
class SalesReportRow:
    sku: str
    name: str
    quantity: int
    amount: Decimal | None


@dataclass(frozen=True)
class SalesReportResult:
    marketplace_id: str
    marketplace_title: str
    date_from: str
    date_to: str
    rows: list[SalesReportRow]
    total_quantity: int
    total_amount: Decimal | None
    missing_amount_count: int
    workbook_bytes: bytes
    filename: str


def _excel_safe_text(value: object) -> str:
    text = str(value or "")
    return "".join(ch for ch in text if ord(ch) >= 32 or ch in "\t\n\r")


def parse_report_dates(date_from: str, date_to: str) -> tuple[str, str, int, int]:
    start_raw = str(date_from or "").strip()
    end_raw = str(date_to or "").strip()
    try:
        start_day = datetime.strptime(start_raw, "%Y-%m-%d")
        end_day = datetime.strptime(end_raw, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("Даты должны быть в формате ГГГГ-ММ-ДД") from exc
    if end_day.date() < start_day.date():
        raise ValueError("Дата окончания не может быть раньше даты начала")
    start_ts = int(start_day.replace(tzinfo=_TZ).timestamp())
    end_ts = int(end_day.replace(hour=23, minute=59, second=59, tzinfo=_TZ).timestamp())
    return start_raw, end_raw, start_ts, end_ts


def build_sales_report(
    orders_repo: WarehouseOrdersRepository,
    catalog_repo: CatalogRepository,
    *,
    marketplace_id: str,
    date_from: str,
    date_to: str,
) -> SalesReportResult:
    source = str(marketplace_id or "").strip().lower()
    if source not in MARKETPLACE_IDS:
        raise ValueError(
            "Маркетплейс должен быть одним из: " + ", ".join(item["title"] for item in MARKETPLACES)
        )
    start_raw, end_raw, start_ts, end_ts = parse_report_dates(date_from, date_to)
    title = marketplace_title(source)

    grouped: dict[str, dict[str, Any]] = {}
    with Session(orders_repo.engine) as session:
        stmt = (
            select(WarehouseOrder, WarehouseOrderLine)
            .join(WarehouseOrderLine, WarehouseOrderLine.order_id == WarehouseOrder.id)
            .where(WarehouseOrder.source == source)
            .where(WarehouseOrder.status != ORDER_CANCELLED)
            .where(WarehouseOrder.created_at_ts >= start_ts)
            .where(WarehouseOrder.created_at_ts <= end_ts)
        )
        for _order, line in session.execute(stmt).all():
            qty = int(line.quantity or 0)
            if qty <= 0:
                continue
            sku = str(line.sku or "").strip()
            if not sku:
                continue
            key = sku.casefold()
            bucket = grouped.get(key)
            if bucket is None:
                bucket = {
                    "sku": sku,
                    "name": str(line.name or "").strip(),
                    "quantity": 0,
                    "amount": Decimal("0.00"),
                    "priced": True,
                }
                grouped[key] = bucket
            bucket["quantity"] += qty
            if not bucket["name"] and str(line.name or "").strip():
                bucket["name"] = str(line.name).strip()
            price = parse_money(line.unit_price)
            if price is None:
                bucket["priced"] = False
            else:
                bucket["amount"] += round_money(price * Decimal(qty))

    catalog_by_sku = catalog_repo.lookup_products_by_skus(list(grouped))
    rows: list[SalesReportRow] = []
    total_qty = 0
    total_amount = Decimal("0.00")
    amount_complete = True
    missing_amount = 0
    for key in sorted(grouped):
        bucket = grouped[key]
        catalog = catalog_by_sku.get(key)
        if catalog:
            sku = str(catalog.get("sku") or bucket["sku"]).strip() or bucket["sku"]
            name = str(catalog.get("name") or "").strip() or bucket["name"]
        else:
            sku = bucket["sku"]
            name = bucket["name"]
        if not name:
            name = MISSING_NAME
        qty = int(bucket["quantity"])
        amount = round_money(bucket["amount"]) if bucket["priced"] else None
        if amount is None:
            missing_amount += 1
            amount_complete = False
        else:
            total_amount += amount
        total_qty += qty
        rows.append(SalesReportRow(sku=sku, name=name, quantity=qty, amount=amount))

    workbook = _build_workbook(
        rows,
        marketplace_title=title,
        date_from=start_raw,
        date_to=end_raw,
        total_quantity=total_qty,
        total_amount=total_amount if rows and amount_complete else None,
    )
    filename = f"otchet_prodazhi_{source}_{start_raw}_{end_raw}.xlsx"
    return SalesReportResult(
        marketplace_id=source,
        marketplace_title=title,
        date_from=start_raw,
        date_to=end_raw,
        rows=rows,
        total_quantity=total_qty,
        total_amount=total_amount if rows and amount_complete else None,
        missing_amount_count=missing_amount,
        workbook_bytes=workbook,
        filename=filename,
    )


def _build_workbook(
    rows: list[SalesReportRow],
    *,
    marketplace_title: str,
    date_from: str,
    date_to: str,
    total_quantity: int,
    total_amount: Decimal | None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    thin = Border(
        left=Side(style="thin", color="D8DEE9"),
        right=Side(style="thin", color="D8DEE9"),
        top=Side(style="thin", color="D8DEE9"),
        bottom=Side(style="thin", color="D8DEE9"),
    )
    money_align = Alignment(horizontal="right")
    ws.append(["Маркетплейс", _excel_safe_text(marketplace_title)])
    ws.append(["Период", f"{date_from} — {date_to}"])
    ws.append([])
    ws.append(list(_HEADERS))
    for col in range(1, len(_HEADERS) + 1):
        cell = ws.cell(row=4, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
    for row in rows:
        amount_val: Any = float(row.amount) if row.amount is not None else None
        ws.append(
            [
                _excel_safe_text(row.sku),
                _excel_safe_text(row.name),
                int(row.quantity),
                amount_val,
            ]
        )
        excel_row = ws.max_row
        if row.amount is not None:
            ws.cell(row=excel_row, column=4).number_format = "#,##0.00"
            ws.cell(row=excel_row, column=4).alignment = money_align
        for col in range(1, 5):
            ws.cell(row=excel_row, column=col).border = thin
    if rows:
        ws.append(
            [
                "Итого",
                "",
                int(total_quantity),
                float(total_amount) if total_amount is not None else None,
            ]
        )
        total_row = ws.max_row
        for col in range(1, 5):
            cell = ws.cell(row=total_row, column=col)
            cell.font = header_font
            cell.border = thin
        if total_amount is not None:
            ws.cell(row=total_row, column=4).number_format = "#,##0.00"
            ws.cell(row=total_row, column=4).alignment = money_align
    ws.column_dimensions[get_column_letter(1)].width = 22
    ws.column_dimensions[get_column_letter(2)].width = 48
    ws.column_dimensions[get_column_letter(3)].width = 16
    ws.column_dimensions[get_column_letter(4)].width = 16
    ws.auto_filter.ref = f"A4:D{max(4, len(rows) + 4)}"
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

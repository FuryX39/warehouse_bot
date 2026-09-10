"""Excel-журнал сырых сканов без сопоставления с каталогом."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.marking.cis import replace_gs_for_excel

_CELL_LIMIT = 32000
_MAX_CODES = 20000


def _safe(value: object) -> str:
    text = replace_gs_for_excel(str(value or ""))
    text = "".join(ch for ch in text if ord(ch) >= 32 or ch in "\t\n\r")
    if len(text) > _CELL_LIMIT:
        return text[:_CELL_LIMIT]
    return text


def _style_header(ws: Worksheet, row: int, columns: int) -> None:
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    thin = Border(
        left=Side(style="thin", color="D8DEE9"),
        right=Side(style="thin", color="D8DEE9"),
        top=Side(style="thin", color="D8DEE9"),
        bottom=Side(style="thin", color="D8DEE9"),
    )
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin


def normalize_scan_codes(raw: object) -> list[str]:
    if isinstance(raw, dict):
        codes = raw.get("codes")
        if codes is None:
            codes = raw.get("scans")
    elif isinstance(raw, list):
        codes = raw
    else:
        codes = None
    if not isinstance(codes, list):
        raise ValueError("codes должен быть массивом")
    out: list[str] = []
    for item in codes:
        if isinstance(item, dict):
            value = item.get("value") or item.get("code") or item.get("raw")
        else:
            value = item
        text = str(value or "").strip(" \t\r\n")
        if not text:
            continue
        out.append(text)
    if not out:
        raise ValueError("Нет отсканированных данных")
    if len(out) > _MAX_CODES:
        raise ValueError(f"Слишком много сканов (макс. {_MAX_CODES})")
    return out


def build_scan_log_export(codes: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сканы"
    headers = ("№", "Скан")
    ws.append(list(headers))
    _style_header(ws, 1, len(headers))
    for idx, code in enumerate(codes, start=1):
        ws.append([idx, _safe(code)])
    ws.column_dimensions[get_column_letter(1)].width = 8
    ws.column_dimensions[get_column_letter(2)].width = 80
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{max(ws.max_row, 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

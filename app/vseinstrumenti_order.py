"""Разбор Excel-подтверждения заказа ВсеИнструменты."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_ORDER_RE = re.compile(r"№\s*(\d+)")
_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
_HEADER_SCAN_FROM = 15
_HEADER_SCAN_TO = 25
_MAX_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class VseinstrumentiOrderLine:
    sku: str
    quantity: int
    excel_name: str
    excel_barcode: str


@dataclass(frozen=True)
class VseinstrumentiOrder:
    order_number: str
    supplier: str
    delivery_date: str
    lines: tuple[VseinstrumentiOrderLine, ...]


def parse_vseinstrumenti_order(content: bytes) -> VseinstrumentiOrder:
    if not content:
        raise ValueError("Файл пустой")
    if len(content) > _MAX_BYTES:
        raise ValueError("Excel слишком большой (макс. 20 МБ)")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось прочитать Excel-файл") from exc
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("В файле нет листов")
        return _parse_sheet(worksheet)
    finally:
        workbook.close()


def _parse_sheet(worksheet: Worksheet) -> VseinstrumentiOrder:
    order_number = _order_number(_cell(worksheet, 1, 1))
    if not order_number:
        raise ValueError("В ячейке A1 нет номера заказа")
    supplier = _labeled_value(worksheet, "поставщик")
    delivery_raw = _labeled_value(worksheet, "ожидаемое время доставки")
    header_row, columns = _header(worksheet)
    grouped: dict[tuple[str, str], VseinstrumentiOrderLine] = {}
    order: list[tuple[str, str]] = []
    for row_idx in range(header_row + 1, (worksheet.max_row or header_row) + 1):
        sku = _cell(worksheet, row_idx, columns["sku"])
        if not sku or sku.casefold() == "итого":
            continue
        barcode = _cell(worksheet, row_idx, columns["barcode"])
        name = _cell(worksheet, row_idx, columns["name"]) if columns["name"] else ""
        quantity = _quantity(
            worksheet.cell(row_idx, columns["confirmed"]).value if columns["confirmed"] else None,
            worksheet.cell(row_idx, columns["ordered"]).value if columns["ordered"] else None,
        )
        if quantity <= 0:
            continue
        key = (sku.casefold(), barcode)
        current = grouped.get(key)
        if current is None:
            grouped[key] = VseinstrumentiOrderLine(
                sku=sku,
                quantity=quantity,
                excel_name=name,
                excel_barcode=barcode,
            )
            order.append(key)
        else:
            grouped[key] = VseinstrumentiOrderLine(
                sku=current.sku,
                quantity=current.quantity + quantity,
                excel_name=current.excel_name or name,
                excel_barcode=current.excel_barcode,
            )
    lines = tuple(grouped[key] for key in order)
    if not lines:
        raise ValueError("В таблице нет товаров с количеством")
    return VseinstrumentiOrder(
        order_number=order_number,
        supplier=supplier,
        delivery_date=_iso_date(delivery_raw),
        lines=lines,
    )


def _header(worksheet: Worksheet) -> tuple[int, dict[str, int]]:
    last = min(worksheet.max_row or _HEADER_SCAN_TO, _HEADER_SCAN_TO)
    for row_idx in range(_HEADER_SCAN_FROM, last + 1):
        labels = {
            _cell(worksheet, row_idx, col).casefold(): col
            for col in range(1, (worksheet.max_column or 1) + 1)
            if _cell(worksheet, row_idx, col)
        }
        sku_col = _find_col(labels, "артикул")
        if sku_col is None:
            continue
        barcode_col = _find_col(labels, "штрихкод") or 2
        return row_idx, {
            "sku": sku_col,
            "barcode": barcode_col,
            "name": _find_col(labels, "наименование") or 0,
            "confirmed": _find_col(labels, "подтверждено") or 0,
            "ordered": _find_col(labels, "заказано") or 0,
        }
    raise ValueError("Не найдена строка заголовка с колонкой «Артикул»")


def _find_col(labels: dict[str, int], needle: str) -> int | None:
    for label, col in labels.items():
        if needle in label:
            return col
    return None


def _labeled_value(worksheet: Worksheet, prefix: str) -> str:
    last = min(worksheet.max_row or 17, 17)
    for row_idx in range(1, last + 1):
        label = _cell(worksheet, row_idx, 1).casefold()
        if not label.startswith(prefix):
            continue
        for col in range(2, (worksheet.max_column or 1) + 1):
            value = _cell(worksheet, row_idx, col)
            if value:
                return value
    return ""


def _order_number(value: str) -> str:
    match = _ORDER_RE.search(value)
    return match.group(1) if match else ""


def _quantity(confirmed: Any, ordered: Any) -> int:
    confirmed_n = _as_int(confirmed)
    if confirmed_n > 0:
        return confirmed_n
    return _as_int(ordered)


def _as_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return 0
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else 0
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return 0
    try:
        number = float(text)
    except ValueError:
        return 0
    if not number.is_integer():
        return 0
    return int(number)


def _iso_date(value: str) -> str:
    match = _DATE_RE.search(value)
    if not match:
        return ""
    return f"{match.group(3)}-{match.group(2)}-{match.group(1)}"


def _cell(worksheet: Worksheet, row_idx: int, col_idx: int) -> str:
    if col_idx <= 0:
        return ""
    return _text(worksheet.cell(row_idx, col_idx).value)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

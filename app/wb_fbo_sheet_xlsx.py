"""Разбор и заполнение Excel FBW: «Товары» и «Шк коробов» в формате кабинета WB."""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

_GOODS_HEADERS = {
    "баркод": "barcode",
    "количество, шт": "qty",
    "количество шт": "qty",
    "количество": "qty",
    "предмет": "subject",
    "артикул поставщика": "sku",
    "бренд": "brand",
    "размер": "size",
    "цвет": "color",
}

_BOX_HEADERS = {
    "баркод товара": "product_barcode",
    "кол-во товаров": "qty",
    "кол-во товара": "qty",
    "количество товаров": "qty",
    "шк короба": "box_id",
    "срок годности": "expiry",
    "шк короба для печати в стороннем сервисе": "package_code",
}


def _norm_header(raw: object) -> str:
    text = str(raw or "").strip().casefold().replace("ё", "е")
    return " ".join(text.split())


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _cell_int(value: object, *, default: int = 0) -> int:
    text = cell_text(value).replace(" ", "").replace(",", ".")
    if not text:
        return default
    try:
        return int(float(text))
    except ValueError:
        return default


def digits_only(raw: object) -> str:
    return "".join(ch for ch in cell_text(raw) if ch.isdigit())


def _map_headers(values: tuple[object, ...], mapping: dict[str, str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, raw in enumerate(values):
        key = _norm_header(raw)
        field = mapping.get(key)
        if field and field not in found:
            found[field] = idx
    return found


def _first_sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    if not content:
        raise ValueError("Файл пустой")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError("Нужен файл Excel (.xlsx)") from exc
    try:
        ws = wb[wb.sheetnames[0]]
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


@dataclass(frozen=True)
class ParsedGoodsRow:
    barcode: str
    qty: int
    sku: str
    name: str
    brand: str
    size: str
    color: str
    subject: str


@dataclass(frozen=True)
class ParsedBoxRow:
    box_id: str
    package_code: str
    product_barcode: str
    qty: int
    expiry: str


def parse_goods_xlsx(content: bytes) -> list[ParsedGoodsRow]:
    rows = _first_sheet_rows(content)
    if not rows:
        raise ValueError("В таблице товаров нет строк")
    cols = _map_headers(rows[0], _GOODS_HEADERS)
    if "barcode" not in cols or "qty" not in cols:
        raise ValueError("В таблице товаров нужны колонки «Баркод» и «Количество, шт»")
    out: list[ParsedGoodsRow] = []
    seen: set[str] = set()
    for raw in rows[1:]:
        barcode = cell_text(raw[cols["barcode"]] if cols["barcode"] < len(raw) else "")
        if not barcode:
            continue
        if barcode.casefold() in seen:
            raise ValueError(f"Баркод «{barcode}» повторяется в таблице товаров")
        seen.add(barcode.casefold())
        qty = _cell_int(raw[cols["qty"]] if cols["qty"] < len(raw) else None)
        if qty <= 0:
            raise ValueError(f"У баркода «{barcode}» количество должно быть больше 0")
        sku = cell_text(raw[cols["sku"]] if "sku" in cols and cols["sku"] < len(raw) else "")
        subject = cell_text(
            raw[cols["subject"]] if "subject" in cols and cols["subject"] < len(raw) else ""
        )
        out.append(
            ParsedGoodsRow(
                barcode=barcode,
                qty=qty,
                sku=sku,
                name=subject or sku or barcode,
                brand=cell_text(
                    raw[cols["brand"]] if "brand" in cols and cols["brand"] < len(raw) else ""
                ),
                size=cell_text(
                    raw[cols["size"]] if "size" in cols and cols["size"] < len(raw) else ""
                ),
                color=cell_text(
                    raw[cols["color"]] if "color" in cols and cols["color"] < len(raw) else ""
                ),
                subject=subject,
            )
        )
    if not out:
        raise ValueError("В таблице товаров нет позиций")
    return out


def parse_boxes_xlsx(content: bytes) -> list[ParsedBoxRow]:
    rows = _first_sheet_rows(content)
    if not rows:
        raise ValueError("В таблице коробов нет строк")
    cols = _map_headers(rows[0], _BOX_HEADERS)
    if "box_id" not in cols or "package_code" not in cols:
        raise ValueError(
            "В таблице коробов нужны колонки «ШК короба» и "
            "«ШК короба для печати в стороннем сервисе»"
        )
    out: list[ParsedBoxRow] = []
    seen_box: set[str] = set()
    seen_pkg: set[str] = set()
    for raw in rows[1:]:
        box_id = cell_text(raw[cols["box_id"]] if cols["box_id"] < len(raw) else "")
        package_code = cell_text(
            raw[cols["package_code"]] if cols["package_code"] < len(raw) else ""
        )
        if not box_id and not package_code:
            continue
        if not box_id or not package_code:
            raise ValueError("У короба должны быть и «ШК короба», и код для печати")
        box_key = digits_only(box_id) or box_id.casefold()
        pkg_key = package_code.casefold()
        if box_key in seen_box:
            raise ValueError(f"ШК короба «{box_id}» повторяется")
        if pkg_key in seen_pkg:
            raise ValueError(f"Код печати короба «{package_code}» повторяется")
        seen_box.add(box_key)
        seen_pkg.add(pkg_key)
        product_barcode = cell_text(
            raw[cols["product_barcode"]]
            if "product_barcode" in cols and cols["product_barcode"] < len(raw)
            else ""
        )
        qty = _cell_int(
            raw[cols["qty"]] if "qty" in cols and cols["qty"] < len(raw) else None
        )
        out.append(
            ParsedBoxRow(
                box_id=box_id,
                package_code=package_code,
                product_barcode=product_barcode,
                qty=qty if product_barcode else 0,
                expiry=cell_text(
                    raw[cols["expiry"]] if "expiry" in cols and cols["expiry"] < len(raw) else ""
                ),
            )
        )
    if not out:
        raise ValueError("В таблице коробов нет грузомест")
    return out


def fill_boxes_xlsx(
    original: bytes,
    assignments: list[dict[str, Any]],
) -> bytes:
    """Пишет «Баркод товара» и «Кол-во товаров» в исходный шаблон WB."""
    if not original:
        raise ValueError("Нет исходной таблицы коробов")
    by_box: dict[str, dict[str, Any]] = {}
    for item in assignments:
        box_id = cell_text(item.get("box_id") or item.get("box_human_id"))
        key = digits_only(box_id) or box_id.casefold()
        if key:
            by_box[key] = item
    try:
        wb = load_workbook(io.BytesIO(original))
    except Exception as exc:
        raise ValueError("Нужен файл Excel (.xlsx)") from exc
    ws = wb[wb.sheetnames[0]]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    cols = _map_headers(tuple(header_row), _BOX_HEADERS)
    if "box_id" not in cols or "product_barcode" not in cols or "qty" not in cols:
        raise ValueError("В шаблоне коробов нет колонок для заполнения")
    barcode_col = cols["product_barcode"] + 1
    qty_col = cols["qty"] + 1
    box_col = cols["box_id"] + 1
    for row_idx in range(2, (ws.max_row or 1) + 1):
        box_id = cell_text(ws.cell(row_idx, box_col).value)
        key = digits_only(box_id) or box_id.casefold()
        item = by_box.get(key)
        barcode_cell = ws.cell(row_idx, barcode_col)
        qty_cell = ws.cell(row_idx, qty_col)
        if item is None:
            continue
        barcode = cell_text(item.get("product_barcode"))
        qty = int(item.get("item_qty") or item.get("quantity") or 0)
        barcode_cell.number_format = "@"
        barcode_cell.value = barcode
        qty_cell.value = qty if barcode and qty > 0 else 0
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def example_column_letter(index: int) -> str:
    return get_column_letter(index + 1)

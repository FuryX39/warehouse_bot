"""Сборка задания ВсеИнструменты и пик строк."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.catalog_repository import CatalogRepository
from app.fbs_packing_service import CIS_REQUIRED_ERROR, resolve_packing_scan
from app.marking.cis import replace_gs_for_excel
from app.marketplace_route_sheets import (
    DEFAULT_ROUTE_SUPPLIER,
    generate_vseinstrumenti_route_sheets_pdf,
    list_route_purchase_statuses,
    normalize_vseinstrumenti_route_sheet_payload,
)
from app.other_marketplace_repository import (
    LINE_DONE,
    PLATFORM_VSEINSTRUMENTI,
    OtherMarketplaceJobRow,
    OtherMarketplaceLineRow,
    OtherMarketplaceRepository,
)
from app.vseinstrumenti_order import VseinstrumentiOrder, parse_vseinstrumenti_order

BARCODE_SOURCE_NOTE = "Печатается штрихкод из столбца «Штрихкод» файла заказа"


def require_purchase_status(value: str) -> str:
    name = str(value or "").strip()
    allowed = {str(row.get("name") or "").strip() for row in list_route_purchase_statuses()}
    if not name or name not in allowed:
        raise ValueError("Укажите статус закупки")
    return name


def barcodes_match(left: str, right: str) -> bool:
    a = str(left or "").strip()
    b = str(right or "").strip()
    if not a or not b:
        return False
    if a.casefold() == b.casefold():
        return True
    a_digits = "".join(ch for ch in a if ch.isdigit())
    b_digits = "".join(ch for ch in b if ch.isdigit())
    return bool(a_digits) and a_digits == b_digits


def create_vseinstrumenti_job(
    *,
    catalog: CatalogRepository,
    repo: OtherMarketplaceRepository,
    content: bytes,
    filename: str,
    transfer_number: str,
    purchase_status: str,
    packer_user_ids: list[int],
    created_by_user_id: int | None,
) -> tuple[OtherMarketplaceJobRow, list[str]]:
    number = str(transfer_number or "").strip()
    status_name = require_purchase_status(purchase_status)
    if not number:
        raise ValueError("Укажите номер перемещения")
    packers = []
    seen: set[int] = set()
    for raw in packer_user_ids:
        try:
            user_id = int(raw)
        except (TypeError, ValueError):
            continue
        if user_id <= 0 or user_id in seen:
            continue
        seen.add(user_id)
        packers.append(user_id)
    if not packers:
        raise ValueError("Выберите хотя бы одного упаковщика")
    order = parse_vseinstrumenti_order(content)
    by_sku = catalog.lookup_products_by_skus([line.sku for line in order.lines])
    warnings: list[str] = []
    payloads: list[dict[str, Any]] = []
    for line in order.lines:
        product = by_sku.get(line.sku.casefold())
        product_id = int(product["id"]) if product else None
        if product is None:
            warnings.append(f"Артикул «{line.sku}» не найден в товарах — в задании название из Excel")
        payloads.append(
            {
                "sku": line.sku,
                "product_id": product_id,
                "product_name": str(product["name"]) if product else line.excel_name,
                "excel_barcode": line.excel_barcode,
                "quantity": line.quantity,
                "require_cis": catalog.product_requires_cis(product_id) if product_id else False,
            }
        )
    job = repo.create_job(
        platform=PLATFORM_VSEINSTRUMENTI,
        order_number=order.order_number,
        transfer_number=number,
        supplier=order.supplier or DEFAULT_ROUTE_SUPPLIER,
        delivery_date=order.delivery_date,
        purchase_status=status_name,
        source_filename=str(filename or "")[:256],
        created_by_user_id=created_by_user_id,
        packer_user_ids=packers,
        lines=payloads,
    )
    return job, warnings


def pick_other_marketplace_line(
    *,
    catalog: CatalogRepository,
    repo: OtherMarketplaceRepository,
    job_id: int,
    raw: str,
    sku: str = "",
    product_id: int | None = None,
) -> dict[str, Any]:
    job = repo.get_job(job_id)
    if job is None:
        raise ValueError("Задание не найдено")
    text = str(raw or "").strip()
    if text:
        pending = [line for line in job.lines if line.status != LINE_DONE]
        excel_line = _match_excel_barcode(pending, text)
        if excel_line is not None:
            if excel_line.require_cis:
                raise ValueError(CIS_REQUIRED_ERROR)
            add_qty = max(0, excel_line.quantity - excel_line.picked_qty)
            updated = repo.record_pick(job.id, excel_line.id, add_qty=add_qty)
            return _pick_payload(repo, job.id, updated, copies=add_qty, mismatch=False)
        resolved = resolve_packing_scan(catalog, text)
        return _apply_resolved(repo, job, resolved.sku, resolved.product_id, text, resolved)
    target_sku = str(sku or "").strip()
    if not target_sku:
        raise ValueError("Пустой штрихкод")
    return _apply_resolved(
        repo,
        job,
        target_sku,
        product_id,
        target_sku,
        _Tap(sku=target_sku, product_id=product_id),
    )


class _Tap:
    def __init__(self, sku: str, product_id: int | None) -> None:
        self.sku = sku
        self.product_id = product_id
        self.is_cis = False
        self.cis_key = ""
        self.cis_raw = ""
        self.cis_gtin = ""


def _apply_resolved(repo, job: OtherMarketplaceJobRow, sku: str, product_id: int | None, raw: str, resolved) -> dict[str, Any]:
    pending = [line for line in job.lines if line.status != LINE_DONE]
    if resolved.is_cis:
        line = _match_line(pending, sku, product_id)
        if line is None:
            raise ValueError("Этого товара нет в задании")
        if not line.require_cis:
            raise ValueError("Для этого товара КИЗ не нужен — пикните штрихкод товара")
        updated = repo.record_pick(
            job.id,
            line.id,
            add_qty=1,
            cis_key=resolved.cis_key,
            cis_raw=resolved.cis_raw,
            cis_gtin=resolved.cis_gtin,
        )
        copies = 1 if updated.picked_qty > line.picked_qty else 0
        return _pick_payload(repo, job.id, updated, copies=copies, mismatch=False)
    line = _match_excel_barcode(pending, raw)
    mismatch = ""
    if line is None:
        line = _match_line(pending, sku, product_id)
        if line is None:
            raise ValueError("Этого товара нет в задании")
        if line.excel_barcode and not barcodes_match(raw, line.excel_barcode) and raw.casefold() != line.sku.casefold():
            mismatch = raw
        if line.require_cis:
            if mismatch:
                repo.record_pick(job.id, line.id, add_qty=0, mismatch_barcode=mismatch)
            raise ValueError(CIS_REQUIRED_ERROR)
    elif line.require_cis:
        raise ValueError(CIS_REQUIRED_ERROR)
    add_qty = max(0, line.quantity - line.picked_qty)
    updated = repo.record_pick(
        job.id,
        line.id,
        add_qty=add_qty,
        mismatch_barcode=mismatch,
    )
    return _pick_payload(repo, job.id, updated, copies=add_qty, mismatch=bool(mismatch))


def _match_excel_barcode(lines: list[OtherMarketplaceLineRow], raw: str) -> OtherMarketplaceLineRow | None:
    for line in lines:
        if barcodes_match(raw, line.excel_barcode):
            return line
    return None


def _match_line(
    lines: list[OtherMarketplaceLineRow],
    sku: str,
    product_id: int | None,
) -> OtherMarketplaceLineRow | None:
    key = str(sku or "").strip().casefold()
    for line in lines:
        if product_id and line.product_id == int(product_id):
            return line
    for line in lines:
        if key and line.sku.casefold() == key:
            return line
    return None


def _pick_payload(repo, job_id: int, line: OtherMarketplaceLineRow, *, copies: int, mismatch: bool) -> dict[str, Any]:
    job = repo.get_job(job_id)
    if job is None:
        raise ValueError("Задание не найдено")
    barcode = line.excel_barcode
    return {
        "job": repo.job_to_dict(job),
        "line_id": line.id,
        "barcode": barcode,
        "barcode_copies": copies if barcode else 0,
        "barcode_note": BARCODE_SOURCE_NOTE if barcode else "В строке заказа нет штрихкода",
        "mismatch": mismatch,
    }


def vseinstrumenti_route_sheet_data(job: OtherMarketplaceJobRow, *, cargo_type: str, cargo_count: int):
    status_name = str(job.purchase_status or "").strip()
    if not status_name:
        raise ValueError("Менеджер не указал статус закупки")
    return normalize_vseinstrumenti_route_sheet_payload(
        {
            "supplier": job.supplier,
            "purchase_number": job.order_number,
            "purchase_status": status_name,
            "delivery_date": job.delivery_date,
            "transfer_number": job.transfer_number,
            "cargo_type": cargo_type,
            "cargo_count": cargo_count,
        }
    )


def build_vseinstrumenti_route_pdf(job: OtherMarketplaceJobRow, *, cargo_type: str, cargo_count: int) -> bytes:
    return generate_vseinstrumenti_route_sheets_pdf(
        vseinstrumenti_route_sheet_data(job, cargo_type=cargo_type, cargo_count=cargo_count)
    )


def attach_catalog_images(catalog: CatalogRepository, payload: dict[str, Any]) -> None:
    buckets: list[dict[str, Any]] = []
    remaining = payload.get("remaining_groups")
    if isinstance(remaining, list):
        buckets.extend(item for item in remaining if isinstance(item, dict))
    lines = payload.get("lines")
    if isinstance(lines, list):
        buckets.extend(item for item in lines if isinstance(item, dict))
    product_ids: list[int] = []
    skus: list[str] = []
    for item in buckets:
        raw_id = item.get("product_id")
        if raw_id not in (None, ""):
            try:
                product_ids.append(int(raw_id))
            except (TypeError, ValueError):
                pass
        sku = str(item.get("sku") or "").strip()
        if sku:
            skus.append(sku)
    by_id = catalog.image_urls_by_product_ids(product_ids) if product_ids else {}
    by_sku = catalog.lookup_products_by_skus(skus) if skus else {}
    for item in buckets:
        url = ""
        raw_id = item.get("product_id")
        if raw_id not in (None, ""):
            try:
                url = by_id.get(int(raw_id), "") or ""
            except (TypeError, ValueError):
                url = ""
        if not url:
            product = by_sku.get(str(item.get("sku") or "").strip().casefold()) or {}
            url = str(product.get("image_url") or "")
        item["image_url"] = url


def build_other_marketplace_marking_xlsx(job: OtherMarketplaceJobRow) -> bytes:
    wb = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF4")

    def write_header(ws, headers: tuple[str, ...]) -> None:
        ws.append(list(headers))
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

    full = wb.active
    full.title = "Все строки"
    write_header(full, ("Заказ", "SKU", "КИЗ"))
    with_cis_rows: list[list[str]] = []
    for line in sorted(job.lines, key=lambda item: (item.seq, item.id)):
        cis_values = [replace_gs_for_excel(item.cis_raw or item.cis_key or "") for item in line.cis]
        slots = max(int(line.quantity), len(cis_values), 1)
        for index in range(slots):
            cis = cis_values[index] if index < len(cis_values) else ""
            full.append([job.order_number, line.sku, cis])
            if cis:
                with_cis_rows.append([job.order_number, line.sku, cis])
    full.column_dimensions["A"].width = 18
    full.column_dimensions["B"].width = 18
    full.column_dimensions["C"].width = 48

    with_cis = wb.create_sheet("С КИЗ")
    write_header(with_cis, ("Заказ", "SKU", "КИЗ"))
    for row in with_cis_rows:
        with_cis.append(row)
    with_cis.column_dimensions["A"].width = 18
    with_cis.column_dimensions["B"].width = 18
    with_cis.column_dimensions["C"].width = 48

    mismatches = wb.create_sheet("Несовпадения ШК")
    write_header(mismatches, ("Заказ", "SKU", "Название", "ШК в поставке", "Отсканированный ШК"))
    for line in job.lines:
        for item in line.mismatches:
            mismatches.append(
                [
                    job.order_number,
                    line.sku,
                    line.product_name,
                    line.excel_barcode,
                    item.scanned_barcode,
                ]
            )
    for col, width in zip("ABCDE", (18, 18, 42, 22, 22)):
        mismatches.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parsed_order_from_bytes(content: bytes) -> VseinstrumentiOrder:
    return parse_vseinstrumenti_order(content)

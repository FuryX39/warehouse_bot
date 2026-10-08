"""Отчёт «Остатки товаров на складах»: сейчас только склады WB (колонка «Склад WB РФ»)."""

from __future__ import annotations

import io
import time
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Callable
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.adapters.base import is_value_configured
from app.catalog_repository import CatalogRepository

ANALYTICS_BASE = "https://seller-analytics-api.wildberries.ru"
SUPPLIES_BASE = "https://supplies-api.wildberries.ru"
WB_RF_WAREHOUSE = "Склад WB РФ"
MISSING_NAME = "Товар не найден в каталоге"
# 2 запланирована, 3 отгрузка разрешена, 6 отгружено на воротах.
IN_TRANSIT_STATUS_IDS = (2, 3, 6)
IN_TRANSIT_FROM_DATE = date(2026, 9, 1)
_TZ = ZoneInfo("Europe/Moscow")
_SUPPLIES_PAGE = 1000
_GOODS_PAGE = 1000
_SUPPLIES_MIN_INTERVAL_SEC = 0.2

SOURCES: tuple[dict[str, Any], ...] = (
    {"id": "own", "title": "Среди своих складов", "implemented": False},
    {"id": "wb", "title": "Склады WB", "implemented": True},
    {"id": "yandex", "title": "Склады Яндекса", "implemented": False},
    {"id": "ozon", "title": "Склады Ozon", "implemented": False},
)
SOURCE_IDS = frozenset(item["id"] for item in SOURCES)
_HEADERS = ("Артикул", "Название", "Количество", "В пути", "Всего")

FetchRawFn = Callable[[str], list[dict[str, Any]]]
FetchInTransitFn = Callable[[str], dict[str, int]]


@dataclass(frozen=True)
class StockByWarehouseRow:
    sku: str
    name: str
    quantity: int
    in_transit: int
    total: int


@dataclass(frozen=True)
class StockByWarehouseResult:
    source_id: str
    source_title: str
    warehouse_title: str
    rows: list[StockByWarehouseRow]
    total_quantity: int
    total_in_transit: int
    total_all: int
    missing_name_count: int
    workbook_bytes: bytes
    filename: str


class _RequestPace:
    def __init__(self, interval_sec: float) -> None:
        self.interval_sec = interval_sec
        self._last = 0.0

    def wait(self) -> None:
        delay = self._last + self.interval_sec - time.monotonic()
        if delay > 0:
            time.sleep(delay)
        self._last = time.monotonic()


def source_meta() -> list[dict[str, Any]]:
    return [dict(item) for item in SOURCES]


def _excel_safe_text(value: object) -> str:
    text = str(value or "")
    return "".join(ch for ch in text if ord(ch) >= 32 or ch in "\t\n\r")


def _qty(value: object) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _as_dict(raw: Any) -> dict[str, Any]:
    return raw if isinstance(raw, dict) else {}


def _as_list(raw: Any) -> list[Any]:
    if isinstance(raw, list):
        return raw
    if isinstance(raw, dict):
        for key in ("data", "report", "reports", "stocks", "items", "supplies", "goods"):
            inner = raw.get(key)
            if isinstance(inner, list):
                return inner
    return []


def _http_error(response: requests.Response) -> requests.HTTPError:
    body = (response.text or "").strip()
    if len(body) > 800:
        body = body[:800] + "…"
    return requests.HTTPError(
        f"{response.status_code} {response.reason} for {response.url} — {body}",
        response=response,
    )


def _json(response: requests.Response) -> Any:
    if response.status_code == 204:
        return None
    if response.status_code == 429:
        raise ValueError("WB принимает этот отчёт не чаще раза в минуту. Повторите через минуту.")
    if not response.ok:
        raise _http_error(response)
    if not (response.content or b"").strip():
        return None
    return response.json()


def _task_id(payload: Any) -> str:
    data = _as_dict(payload)
    nested = _as_dict(data.get("data"))
    for blob in (nested, data):
        for key in ("taskId", "task_id", "id"):
            value = str(blob.get(key) or "").strip()
            if value:
                return value
    raise RuntimeError("WB не вернул идентификатор отчёта остатков")


def _task_status(payload: Any) -> str:
    data = _as_dict(payload)
    nested = _as_dict(data.get("data"))
    for blob in (nested, data):
        value = str(blob.get("status") or blob.get("Status") or "").strip().lower()
        if value:
            return value
    return ""


def _positive_int(raw: object) -> int | None:
    if raw in (None, "", 0, "0"):
        return None
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _parse_wb_date(raw: object) -> date | None:
    text = str(raw or "").strip()
    if not text:
        return None
    if text.endswith("Z"):
        text = text[:-1] + "+00:00"
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        try:
            parsed = datetime.fromisoformat(text[:10])
        except ValueError:
            return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=_TZ)
    return parsed.astimezone(_TZ).date()


def supply_matches_in_transit_cutoff(item: dict[str, Any]) -> bool:
    created = _parse_wb_date(item.get("createDate"))
    if created is not None:
        return created >= IN_TRANSIT_FROM_DATE
    planned = _parse_wb_date(item.get("supplyDate"))
    if planned is not None:
        return planned >= IN_TRANSIT_FROM_DATE
    return False


def _in_transit_list_body() -> dict[str, Any]:
    till = datetime.now(_TZ).date()
    if till < IN_TRANSIT_FROM_DATE:
        till = IN_TRANSIT_FROM_DATE
    return {
        "statusIDs": list(IN_TRANSIT_STATUS_IDS),
        "dates": [
            {
                "type": "createDate",
                "from": f"{IN_TRANSIT_FROM_DATE.isoformat()}T00:00:00+03:00",
                "till": f"{till.isoformat()}T23:59:59+03:00",
            }
        ],
    }


def _supply_fetch_key(item: dict[str, Any]) -> tuple[int, bool] | None:
    supply_id = _positive_int(item.get("supplyID"))
    if supply_id:
        return supply_id, False
    preorder_id = _positive_int(item.get("preorderID"))
    if preorder_id:
        return preorder_id, True
    return None


def _supplies_json(
    method: str,
    url: str,
    *,
    headers: dict[str, str],
    json: Any = None,
    params: dict[str, Any] | None = None,
    pace: _RequestPace | None = None,
    retry_429: int = 4,
) -> Any:
    for attempt in range(retry_429 + 1):
        if pace is not None:
            pace.wait()
        response = requests.request(
            method,
            url,
            headers=headers,
            json=json,
            params=params,
            timeout=90,
        )
        if response.status_code == 429:
            time.sleep(2 * (attempt + 1))
            continue
        if not response.ok:
            raise _http_error(response)
        if response.status_code == 204 or not (response.content or b"").strip():
            return None
        return response.json()
    raise ValueError("WB ограничивает частоту запросов к поставкам. Повторите через минуту.")


def fetch_wb_warehouse_remains(token: str, *, timeout_sec: int = 180) -> list[dict[str, Any]]:
    if not is_value_configured(token):
        raise ValueError("Не задан WB_API_TOKEN")
    headers = {"Authorization": token}
    params = {
        "locale": "ru",
        "groupByBrand": "true",
        "groupBySubject": "true",
        "groupBySa": "true",
        "groupByNm": "true",
        "groupByBarcode": "true",
        "groupBySize": "true",
    }
    try:
        created = requests.get(
            f"{ANALYTICS_BASE}/api/v1/warehouse_remains",
            headers=headers,
            params=params,
            timeout=60,
        )
        task_id = _task_id(_json(created))
    except requests.HTTPError as exc:
        status = exc.response.status_code if exc.response is not None else 0
        if status in (401, 403):
            raise ValueError(
                "Нет доступа к аналитике WB. Нужен токен с категорией «Аналитика»."
            ) from exc
        raise RuntimeError(str(exc)) from exc
    status_url = f"{ANALYTICS_BASE}/api/v1/warehouse_remains/tasks/{task_id}/status"
    deadline = time.monotonic() + timeout_sec
    last = ""
    try:
        while time.monotonic() < deadline:
            status_resp = requests.get(status_url, headers=headers, timeout=60)
            if status_resp.status_code == 429:
                time.sleep(6)
                continue
            last = _task_status(_json(status_resp))
            if last in {"done", "success", "ready"}:
                break
            if last in {"canceled", "cancelled", "purged", "error", "failed"}:
                raise RuntimeError(f"Отчёт WB завершился со статусом {last}")
            time.sleep(6)
        else:
            raise TimeoutError(f"Отчёт WB не готов за {timeout_sec} с")
        download = requests.get(
            f"{ANALYTICS_BASE}/api/v1/warehouse_remains/tasks/{task_id}/download",
            headers=headers,
            timeout=120,
        )
        payload = _json(download)
    except requests.HTTPError as exc:
        raise RuntimeError(str(exc)) from exc
    return [item for item in _as_list(payload) if isinstance(item, dict)]


def _list_in_transit_supplies(token: str, pace: _RequestPace) -> list[dict[str, Any]]:
    headers = {"Authorization": token, "Content-Type": "application/json"}
    offset = 0
    out: list[dict[str, Any]] = []
    while True:
        payload = _supplies_json(
            "POST",
            f"{SUPPLIES_BASE}/api/v1/supplies",
            headers=headers,
            params={"limit": _SUPPLIES_PAGE, "offset": offset},
            json=_in_transit_list_body(),
            pace=pace,
        )
        items = [item for item in _as_list(payload) if isinstance(item, dict)]
        if not items:
            break
        out.extend(item for item in items if supply_matches_in_transit_cutoff(item))
        if len(items) < _SUPPLIES_PAGE:
            break
        offset += len(items)
    return out


def _fetch_supply_goods(
    token: str,
    supply_id: int,
    *,
    is_preorder: bool,
    pace: _RequestPace,
) -> list[dict[str, Any]]:
    headers = {"Authorization": token}
    offset = 0
    out: list[dict[str, Any]] = []
    while True:
        params: dict[str, Any] = {"limit": _GOODS_PAGE, "offset": offset}
        if is_preorder:
            params["isPreorderID"] = "true"
        payload = _supplies_json(
            "GET",
            f"{SUPPLIES_BASE}/api/v1/supplies/{supply_id}/goods",
            headers=headers,
            params=params,
            pace=pace,
        )
        items = [item for item in _as_list(payload) if isinstance(item, dict)]
        if not items:
            break
        out.extend(items)
        if len(items) < _GOODS_PAGE:
            break
        offset += len(items)
    return out


def aggregate_in_transit_quantities(goods_rows: list[dict[str, Any]]) -> dict[str, int]:
    by_key: dict[str, tuple[str, int]] = {}
    for row in goods_rows:
        sku = str(row.get("vendorCode") or "").strip()
        if not sku:
            continue
        qty = _qty(row.get("quantity"))
        if qty <= 0:
            continue
        key = sku.casefold()
        prev_sku, prev_qty = by_key.get(key, (sku, 0))
        by_key[key] = (prev_sku, prev_qty + qty)
    return {sku: qty for sku, qty in by_key.values()}


def fetch_wb_in_transit_quantities(token: str) -> dict[str, int]:
    if not is_value_configured(token):
        raise ValueError("Не задан WB_API_TOKEN")
    pace = _RequestPace(_SUPPLIES_MIN_INTERVAL_SEC)
    try:
        supplies = _list_in_transit_supplies(token, pace)
    except requests.HTTPError as exc:
        status = exc.response.status_code if exc.response is not None else 0
        if status in (401, 403):
            raise ValueError(
                "Нет доступа к поставкам WB. Нужен токен с категорией «Поставки»."
            ) from exc
        raise RuntimeError(str(exc)) from exc
    by_key: dict[str, tuple[str, int]] = {}
    for item in supplies:
        fetch_key = _supply_fetch_key(item)
        if fetch_key is None:
            continue
        supply_id, is_preorder = fetch_key
        try:
            goods = _fetch_supply_goods(token, supply_id, is_preorder=is_preorder, pace=pace)
        except requests.HTTPError as exc:
            status = exc.response.status_code if exc.response is not None else 0
            if status == 404:
                continue
            if status in (401, 403):
                raise ValueError(
                    "Нет доступа к поставкам WB. Нужен токен с категорией «Поставки»."
                ) from exc
            raise RuntimeError(str(exc)) from exc
        for sku, qty in aggregate_in_transit_quantities(goods).items():
            key = sku.casefold()
            prev_sku, prev_qty = by_key.get(key, (sku, 0))
            by_key[key] = (prev_sku, prev_qty + qty)
    return {sku: qty for sku, qty in by_key.values()}


def aggregate_wb_rf_quantities(raw_rows: list[dict[str, Any]]) -> dict[str, int]:
    wanted = WB_RF_WAREHOUSE.casefold()
    by_key: dict[str, tuple[str, int]] = {}
    for row in raw_rows:
        sku = str(row.get("vendorCode") or row.get("supplierArticle") or "").strip()
        if not sku:
            continue
        qty = 0
        warehouses = row.get("warehouses")
        if isinstance(warehouses, list):
            for item in warehouses:
                if not isinstance(item, dict):
                    continue
                name = str(item.get("warehouseName") or "").strip()
                if name.casefold() == wanted:
                    qty += _qty(item.get("quantity"))
        else:
            name = str(row.get("warehouseName") or "").strip()
            if name.casefold() == wanted:
                qty = _qty(
                    row.get("quantity") if row.get("quantity") is not None else row.get("quantityFull")
                )
        if qty <= 0:
            continue
        key = sku.casefold()
        prev_sku, prev_qty = by_key.get(key, (sku, 0))
        by_key[key] = (prev_sku, prev_qty + qty)
    return {sku: qty for sku, qty in by_key.values()}


def build_wb_stock_rows(
    raw_rows: list[dict[str, Any]],
    catalog_by_sku: dict[str, dict[str, Any]],
    in_transit: dict[str, int] | None = None,
) -> list[StockByWarehouseRow]:
    quantities = aggregate_wb_rf_quantities(raw_rows)
    transit = dict(in_transit or {})
    stock_by_key = {sku.casefold(): (sku, qty) for sku, qty in quantities.items()}
    transit_by_key = {sku.casefold(): (sku, qty) for sku, qty in transit.items()}
    rows: list[StockByWarehouseRow] = []
    for key in sorted(set(stock_by_key) | set(transit_by_key)):
        sku = stock_by_key.get(key, transit_by_key.get(key, (key, 0)))[0]
        stock_qty = stock_by_key.get(key, (sku, 0))[1]
        transit_qty = transit_by_key.get(key, (sku, 0))[1]
        if stock_qty <= 0 and transit_qty <= 0:
            continue
        catalog = catalog_by_sku.get(key)
        if catalog:
            display_sku = str(catalog.get("sku") or sku).strip() or sku
            name = str(catalog.get("name") or "").strip() or MISSING_NAME
        else:
            display_sku = sku
            name = MISSING_NAME
        rows.append(
            StockByWarehouseRow(
                sku=display_sku,
                name=name,
                quantity=int(stock_qty),
                in_transit=int(transit_qty),
                total=int(stock_qty) + int(transit_qty),
            )
        )
    return rows


def build_stock_workbook(
    rows: list[StockByWarehouseRow],
    *,
    source_title: str,
    warehouse_title: str,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    thin = Border(
        left=Side(style="thin", color="D8DEE9"),
        right=Side(style="thin", color="D8DEE9"),
        top=Side(style="thin", color="D8DEE9"),
        bottom=Side(style="thin", color="D8DEE9"),
    )
    ncols = len(_HEADERS)
    ws.append(["Источник", _excel_safe_text(source_title)])
    ws.append(["Склад", _excel_safe_text(warehouse_title)])
    ws.append([])
    ws.append(list(_HEADERS))
    for col in range(1, ncols + 1):
        cell = ws.cell(row=4, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
    total_quantity = 0
    total_in_transit = 0
    total_all = 0
    for row in rows:
        ws.append(
            [
                _excel_safe_text(row.sku),
                _excel_safe_text(row.name),
                int(row.quantity),
                int(row.in_transit),
                int(row.total),
            ]
        )
        excel_row = ws.max_row
        for col in range(1, ncols + 1):
            ws.cell(row=excel_row, column=col).border = thin
        total_quantity += int(row.quantity)
        total_in_transit += int(row.in_transit)
        total_all += int(row.total)
    if rows:
        ws.append(["Итого", "", total_quantity, total_in_transit, total_all])
        total_row = ws.max_row
        for col in range(1, ncols + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.font = header_font
            cell.border = thin
        for col in range(3, ncols + 1):
            ws.cell(row=total_row, column=col).alignment = Alignment(horizontal="right")
    ws.column_dimensions[get_column_letter(1)].width = 22
    ws.column_dimensions[get_column_letter(2)].width = 56
    for col in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    last_data_row = max(4, len(rows) + 4)
    ws.auto_filter.ref = f"A4:{get_column_letter(ncols)}{last_data_row}"
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_stock_by_warehouse_report(
    catalog_repo: CatalogRepository,
    *,
    source: str,
    wb_api_token: str = "",
    fetch_wb_raw: FetchRawFn | None = None,
    fetch_wb_in_transit: FetchInTransitFn | None = None,
) -> StockByWarehouseResult:
    source_id = str(source or "").strip().lower()
    if source_id not in SOURCE_IDS:
        raise ValueError(
            "Источник должен быть одним из: " + ", ".join(item["title"] for item in SOURCES)
        )
    item = next(row for row in SOURCES if row["id"] == source_id)
    if not item["implemented"]:
        raise ValueError(f"Выгрузка «{item['title']}» пока не реализована")
    if source_id != "wb":
        raise ValueError(f"Выгрузка «{item['title']}» пока не реализована")
    fetcher = fetch_wb_raw or fetch_wb_warehouse_remains
    raw_rows = fetcher(wb_api_token)
    if fetch_wb_in_transit is not None:
        in_transit = fetch_wb_in_transit(wb_api_token)
    elif fetch_wb_raw is not None:
        in_transit = {}
    else:
        in_transit = fetch_wb_in_transit_quantities(wb_api_token)
    quantities = aggregate_wb_rf_quantities(raw_rows)
    catalog_keys = list({*quantities, *in_transit})
    catalog_by_sku = catalog_repo.lookup_products_by_skus(catalog_keys)
    rows = build_wb_stock_rows(raw_rows, catalog_by_sku, in_transit)
    missing = sum(1 for row in rows if row.name == MISSING_NAME)
    title = str(item["title"])
    workbook = build_stock_workbook(rows, source_title=title, warehouse_title=WB_RF_WAREHOUSE)
    return StockByWarehouseResult(
        source_id=source_id,
        source_title=title,
        warehouse_title=WB_RF_WAREHOUSE,
        rows=rows,
        total_quantity=sum(row.quantity for row in rows),
        total_in_transit=sum(row.in_transit for row in rows),
        total_all=sum(row.total for row in rows),
        missing_name_count=missing,
        workbook_bytes=workbook,
        filename="ostatki_sklady_wb.xlsx",
    )

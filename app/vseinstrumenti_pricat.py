"""Формирование PRICAT ВсеИнструменты из выгрузки свободных остатков."""

from __future__ import annotations

from datetime import date
from io import BytesIO
import math
from pathlib import Path
import xml.etree.ElementTree as ET
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_TEMPLATE_PATH = Path(__file__).resolve().parent / "assets" / "vseinstrumenti_pricat.xlsx"
_SHEET_XML = "xl/worksheets/sheet1.xml"
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _article(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_name(value: Any) -> str:
    return " ".join(str(value or "").strip().casefold().split())


_THOUSAND_SEPARATORS = (" ", "\u00a0", "\u202f", "\u2009", "\u2007", "'", "\u2019")


def _quantity(value: Any, row: int) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        raise ValueError(f"Строка {row}: некорректное количество «{value}»")
    if isinstance(value, int):
        if value < 0:
            raise ValueError(f"Строка {row}: количество должно быть не меньше нуля")
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or value < 0:
            raise ValueError(f"Строка {row}: количество должно быть не меньше нуля")
        return math.floor(value)

    text = str(value).strip()
    if not text:
        return 0
    for sep in _THOUSAND_SEPARATORS:
        text = text.replace(sep, "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Строка {row}: некорректное количество «{value}»") from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError(f"Строка {row}: количество должно быть не меньше нуля")
    return math.floor(number)


def list_vseinstrumenti_pricat_skus(template_bytes: bytes | None = None) -> set[str]:
    template = template_bytes if template_bytes is not None else _TEMPLATE_PATH.read_bytes()
    book = load_workbook(BytesIO(template), data_only=True, read_only=True)
    try:
        if "Лист 1" not in book.sheetnames:
            raise ValueError("В PRICAT отсутствует лист «Лист 1»")
        sheet = book["Лист 1"]
        skus: set[str] = set()
        for row in sheet.iter_rows(min_row=13, min_col=5, max_col=5, values_only=True):
            sku = _article(row[0]).casefold()
            if sku:
                skus.add(sku)
        return skus
    finally:
        book.close()


def quantity_template_names(
    products: list[dict[str, Any]],
    *,
    template_bytes: bytes | None = None,
) -> list[str]:
    pricat_skus = list_vseinstrumenti_pricat_skus(template_bytes)
    names: list[str] = []
    seen: set[str] = set()
    for product in sorted(
        products,
        key=lambda item: (
            _normalize_name(item.get("name")),
            _article(item.get("sku")).casefold(),
        ),
    ):
        if bool(product.get("is_kit")):
            continue
        sku = _article(product.get("sku")).casefold()
        if sku not in pricat_skus:
            continue
        name = str(product.get("name") or "").strip()
        key = _normalize_name(name)
        if not key or key in seen:
            continue
        seen.add(key)
        names.append(name)
    return names


def build_vseinstrumenti_quantity_template(product_names: list[str] | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Количество"
    sheet.append(["Название", "Количество"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315D8A")
    seen: set[str] = set()
    for raw_name in product_names or []:
        name = str(raw_name or "").strip()
        key = _normalize_name(name)
        if not name or key in seen:
            continue
        seen.add(key)
        sheet.append([name, None])
    sheet.freeze_panes = "A2"
    last_row = max(1, sheet.max_row)
    sheet.auto_filter.ref = f"A1:B{last_row}"
    sheet.column_dimensions[get_column_letter(1)].width = 72
    sheet.column_dimensions[get_column_letter(2)].width = 16
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _patch_pricat_xml(template: bytes, values: dict[str, int | str]) -> bytes:
    """Меняет только значения ячеек, сохраняя неподдерживаемые openpyxl расширения XLSX."""
    source = BytesIO(template)
    output = BytesIO()
    with ZipFile(source, "r") as src, ZipFile(output, "w", ZIP_DEFLATED) as dst:
        if _SHEET_XML not in src.namelist():
            raise ValueError("В шаблоне PRICAT не найден XML листа «Лист 1»")
        tree = ET.fromstring(src.read(_SHEET_XML))
        cells = {
            cell.get("r"): cell
            for cell in tree.iter(f"{{{_MAIN_NS}}}c")
            if cell.get("r")
        }
        for coordinate, value in values.items():
            cell = cells.get(coordinate)
            if cell is None:
                raise ValueError(f"В шаблоне PRICAT отсутствует ячейка {coordinate}")
            for child in list(cell):
                if child.tag in {f"{{{_MAIN_NS}}}v", f"{{{_MAIN_NS}}}is", f"{{{_MAIN_NS}}}f"}:
                    cell.remove(child)
            if isinstance(value, str):
                cell.set("t", "inlineStr")
                inline = ET.SubElement(cell, f"{{{_MAIN_NS}}}is")
                text = ET.SubElement(inline, f"{{{_MAIN_NS}}}t")
                text.text = value
            else:
                cell.attrib.pop("t", None)
                node = ET.SubElement(cell, f"{{{_MAIN_NS}}}v")
                node.text = str(value)
        for item in src.infolist():
            data = (
                ET.tostring(tree, encoding="utf-8", xml_declaration=True)
                if item.filename == _SHEET_XML
                else src.read(item.filename)
            )
            dst.writestr(item, data)
    return output.getvalue()


def _read_named_quantities(data: bytes) -> tuple[dict[str, int], dict[str, str]]:
    try:
        book = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("Не удалось прочитать файл количества Excel") from exc
    sheet = book.active
    quantities: dict[str, int] = {}
    display_names: dict[str, str] = {}
    for row in range(1, sheet.max_row + 1):
        raw_name = sheet.cell(row, 1).value
        name = _normalize_name(raw_name)
        if not name:
            continue
        if name in {"название", "наименование", "name"}:
            continue
        value = _quantity(sheet.cell(row, 2).value, row)
        display_names.setdefault(name, str(raw_name).strip())
        if name in quantities:
            quantities[name] = min(quantities[name], value)
        else:
            quantities[name] = value
    book.close()
    return quantities, display_names


def _catalog_indexes(
    products: list[dict[str, Any]],
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, dict[str, Any]]]:
    by_name: dict[str, list[dict[str, Any]]] = {}
    by_sku: dict[str, dict[str, Any]] = {}
    for product in products:
        name = _normalize_name(product.get("name"))
        sku = _article(product.get("sku")).casefold()
        if name:
            by_name.setdefault(name, []).append(product)
        if sku:
            if sku in by_sku:
                raise ValueError(f"В каталоге повторяется артикул «{product.get('sku')}»")
            by_sku[sku] = product
    return by_name, by_sku


def _leaf_bom(
    sku: str,
    products_by_sku: dict[str, dict[str, Any]],
    *,
    trail: tuple[str, ...] = (),
) -> dict[str, int]:
    key = sku.casefold()
    if key in trail:
        raise ValueError(f"Циклический состав комплекта «{sku}»")
    product = products_by_sku.get(key)
    if product is None or not bool(product.get("is_kit")):
        return {key: 1}
    components = product.get("components") or []
    if not components:
        raise ValueError(f"У комплекта «{product.get('sku')}» не указан состав")
    result: dict[str, int] = {}
    for component in components:
        component_sku = _article(component.get("sku")).casefold()
        try:
            component_qty = max(1, int(component.get("quantity") or 1))
        except (TypeError, ValueError):
            component_qty = 1
        if not component_sku or component_sku not in products_by_sku:
            raise ValueError(f"У комплекта «{product.get('sku')}» не найден компонент")
        nested = _leaf_bom(
            component_sku,
            products_by_sku,
            trail=(*trail, key),
        )
        for leaf_sku, leaf_qty in nested.items():
            result[leaf_sku] = result.get(leaf_sku, 0) + leaf_qty * component_qty
    return result


def _format_date(value: date, prefix: str) -> str:
    return f"{prefix} {value:%d.%m.%Y}"


def _required_party(party: dict[str, Any], title: str) -> dict[str, str]:
    fields = {
        "name": str(party.get("full_name") or "").strip(),
        "inn": str(party.get("inn") or "").strip(),
        "kpp": str(party.get("kpp") or "").strip(),
        "gln": str(party.get("gln") or "").strip(),
    }
    missing = [label for key, label in (("name", "наименование"), ("inn", "ИНН"), ("kpp", "КПП"), ("gln", "GLN")) if not fields[key]]
    if missing:
        raise ValueError(f"{title}: не заполнены {', '.join(missing)}")
    return fields


def build_vseinstrumenti_pricat(
    quantity_bytes: bytes,
    *,
    catalog_products: list[dict[str, Any]],
    buyer: dict[str, Any],
    supplier: dict[str, Any],
    header: dict[str, Any],
    template_bytes: bytes | None = None,
    actual_date: date | None = None,
) -> tuple[bytes, dict[str, int]]:
    """Полный PRICAT на основе мастер-шаблона и количества по названиям товаров."""
    named_quantities, display_names = _read_named_quantities(quantity_bytes)
    by_name, by_sku = _catalog_indexes(catalog_products)
    buyer_fields = _required_party(buyer, "Покупатель")
    supplier_fields = _required_party(supplier, "Поставщик")

    try:
        template = template_bytes if template_bytes is not None else _TEMPLATE_PATH.read_bytes()
        template_book = load_workbook(BytesIO(template))
    except Exception as exc:
        raise ValueError("Не удалось прочитать шаблон PRICAT Excel") from exc
    if "Лист 1" not in template_book.sheetnames:
        raise ValueError("В PRICAT отсутствует лист «Лист 1»")
    sheet = template_book["Лист 1"]
    if _article(sheet["E12"].value).casefold() != "артикул":
        raise ValueError("В PRICAT не найден столбец «Артикул» в E12")
    if _article(sheet["AB12"].value).casefold() != "остаток на складе":
        raise ValueError("В PRICAT не найден столбец «Остаток на складе» в AB12")

    pricat_rows: dict[str, int] = {}
    for row in range(13, sheet.max_row + 1):
        sku = _article(sheet.cell(row, 5).value).casefold()
        if not sku:
            continue
        if sku in pricat_rows:
            raise ValueError(f"В PRICAT повторяется артикул «{sheet.cell(row, 5).value}»")
        pricat_rows[sku] = row

    ordinary_quantities: dict[str, int] = {}
    unknown_names: list[str] = []
    ambiguous_names: list[str] = []
    outside_pricat: list[str] = []
    ignored_kit_inputs = 0
    for normalized_name, quantity in named_quantities.items():
        candidates = by_name.get(normalized_name) or []
        if not candidates:
            unknown_names.append(display_names.get(normalized_name, normalized_name))
            continue
        if len(candidates) > 1:
            ambiguous_names.append(str(candidates[0].get("name") or normalized_name))
            continue
        product = candidates[0]
        sku = _article(product.get("sku")).casefold()
        if sku not in pricat_rows:
            outside_pricat.append(str(product.get("name") or normalized_name))
            continue
        if bool(product.get("is_kit")):
            ignored_kit_inputs += 1
            continue
        ordinary_quantities[sku] = quantity

    final_quantities: dict[str, int] = {}
    kits_calculated = 0
    for sku in pricat_rows:
        product = by_sku.get(sku)
        if product and bool(product.get("is_kit")):
            bom = _leaf_bom(sku, by_sku)
            candidates = [
                ordinary_quantities.get(component_sku, 0) // required
                for component_sku, required in bom.items()
            ]
            final_quantities[sku] = min(candidates) if candidates else 0
            kits_calculated += 1
        else:
            final_quantities[sku] = ordinary_quantities.get(sku, 0)

    updates: dict[str, int | str] = {}
    current = actual_date or date.today()
    document_date = header.get("document_date") or current
    contract_date = header.get("contract_date") or current
    valid_from = header.get("valid_from") or current
    valid_to = header.get("valid_to") or current
    updates.update(
        {
            "C1": str(header.get("document_name") or "").strip(),
            "D1": _format_date(document_date, "от"),
            "C2": str(header.get("contract_number") or "").strip(),
            "D2": _format_date(contract_date, "от"),
            "C3": str(header.get("price_list_type") or "Основной").strip(),
            "C4": _format_date(valid_from, "с"),
            "D4": _format_date(valid_to, "по"),
            "C5": str(header.get("internal_comment") or "").strip(),
            "C6": str(header.get("buyer_message") or f"Актуальные остатки {current:%d.%m}").strip(),
            "C9": buyer_fields["name"],
            "D9": buyer_fields["inn"],
            "E9": buyer_fields["kpp"],
            "F9": buyer_fields["gln"],
            "C10": supplier_fields["name"],
            "D10": supplier_fields["inn"],
            "E10": supplier_fields["kpp"],
            "F10": supplier_fields["gln"],
        }
    )
    for sku, row in pricat_rows.items():
        updates[f"AB{row}"] = final_quantities.get(sku, 0)

    template_book.close()
    return _patch_pricat_xml(template, updates), {
        "source_names": len(named_quantities),
        "matched_names": len(ordinary_quantities),
        "skipped_names": len(unknown_names) + len(ambiguous_names) + len(outside_pricat),
        "unknown_names": len(unknown_names),
        "ambiguous_names": len(ambiguous_names),
        "outside_pricat": len(outside_pricat),
        "rows_written": len(pricat_rows),
        "kits_calculated": kits_calculated,
        "ignored_kit_inputs": ignored_kit_inputs,
    }

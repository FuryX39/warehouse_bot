from __future__ import annotations

from datetime import date
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

import pytest

from app.vseinstrumenti_pricat import (
    build_vseinstrumenti_pricat,
    build_vseinstrumenti_quantity_template,
)
from app.crm_repository import CrmRepository


def _quantities(rows=None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["Название", "Количество"])
    for row in rows or []:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _pricat() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист 1"
    for coordinate in (
        "C1", "D1", "C2", "D2", "C3", "C4", "D4", "C5", "C6",
        "C9", "D9", "E9", "F9", "C10", "D10", "E10", "F10",
    ):
        sheet[coordinate] = "old"
    sheet["E12"] = "Артикул"
    sheet["AB12"] = "Остаток на складе"
    sheet["E13"] = "A-1"
    sheet["AB13"] = 999
    sheet["E14"] = "A-2"
    sheet["AB14"] = 999
    sheet["E15"] = "KIT-1"
    sheet["AB15"] = 777
    sheet["E16"] = "ONLY-PRICAT"
    sheet["AB16"] = 777
    sheet["E17"] = "INNER"
    sheet["AB17"] = 777
    sheet["E18"] = "OUTER"
    sheet["AB18"] = 777
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _products() -> list[dict]:
    return [
        {"name": "Первый товар", "sku": "A-1", "is_kit": False, "components": []},
        {"name": "Второй товар", "sku": "A-2", "is_kit": False, "components": []},
        {
            "name": "Комплект",
            "sku": "KIT-1",
            "is_kit": True,
            "components": [
                {"sku": "A-1", "quantity": 2},
                {"sku": "A-2", "quantity": 1},
            ],
        },
        {"name": "Только PRICAT", "sku": "ONLY-PRICAT", "is_kit": False, "components": []},
        {
            "name": "Внутренний комплект",
            "sku": "INNER",
            "is_kit": True,
            "components": [{"sku": "A-1", "quantity": 3}],
        },
        {
            "name": "Внешний комплект",
            "sku": "OUTER",
            "is_kit": True,
            "components": [
                {"sku": "INNER", "quantity": 2},
                {"sku": "A-2", "quantity": 1},
            ],
        },
    ]


def _party(name: str) -> dict:
    return {"full_name": name, "inn": "123", "kpp": "456", "gln": "4600000000000"}


def _header() -> dict:
    return {
        "document_name": "Остатки №7",
        "document_date": date(2026, 9, 24),
        "contract_number": "Д-15",
        "contract_date": date(2026, 1, 10),
        "price_list_type": "Основной",
        "valid_from": date(2026, 9, 25),
        "valid_to": date(2026, 12, 31),
        "internal_comment": "Внутренний",
        "buyer_message": "",
    }


def test_build_vseinstrumenti_pricat_with_kits_and_full_header() -> None:
    result, stats = build_vseinstrumenti_pricat(
        _quantities([
            ["Первый   товар", 8],
            ["первый товар", 6],
            ["Второй товар", 5],
            ["Комплект", 999],
        ]),
        catalog_products=_products(),
        buyer=_party("Покупатель"),
        supplier=_party("Поставщик"),
        header=_header(),
        template_bytes=_pricat(),
        actual_date=date(2026, 9, 24),
    )
    sheet = load_workbook(BytesIO(result), data_only=True)["Лист 1"]

    assert sheet["AB13"].value == 6
    assert sheet["AB14"].value == 5
    assert sheet["AB15"].value == 3
    assert sheet["AB16"].value == 0
    assert sheet["AB17"].value == 2
    assert sheet["AB18"].value == 1
    assert sheet["C1"].value == "Остатки №7"
    assert sheet["D1"].value == "от 24.09.2026"
    assert sheet["C2"].value == "Д-15"
    assert sheet["D2"].value == "от 10.01.2026"
    assert sheet["C4"].value == "с 25.09.2026"
    assert sheet["D4"].value == "по 31.12.2026"
    assert sheet["C6"].value == "Актуальные остатки 24.09"
    assert sheet["C9"].value == "Покупатель"
    assert sheet["C10"].value == "Поставщик"
    assert stats["rows_written"] == 6
    assert stats["kits_calculated"] == 3
    assert stats["ignored_kit_inputs"] == 1


def test_unknown_name_is_rejected() -> None:
    with pytest.raises(ValueError, match="Не найдены в каталоге"):
        build_vseinstrumenti_pricat(
            _quantities([["Неизвестный", 2]]),
            catalog_products=_products(),
            buyer=_party("Покупатель"),
            supplier=_party("Поставщик"),
            header=_header(),
            template_bytes=_pricat(),
        )


def test_ambiguous_catalog_name_is_rejected() -> None:
    products = _products() + [
        {"name": "Первый товар", "sku": "A-3", "is_kit": False, "components": []}
    ]
    with pytest.raises(ValueError, match="Неоднозначные названия"):
        build_vseinstrumenti_pricat(
            _quantities([["Первый товар", 2]]),
            catalog_products=products,
            buyer=_party("Покупатель"),
            supplier=_party("Поставщик"),
            header=_header(),
            template_bytes=_pricat(),
        )


def test_catalog_product_outside_pricat_is_rejected() -> None:
    products = _products() + [
        {"name": "Лишний товар", "sku": "A-3", "is_kit": False, "components": []}
    ]
    with pytest.raises(ValueError, match="Нет в базовом PRICAT"):
        build_vseinstrumenti_pricat(
            _quantities([["Лишний товар", 2]]),
            catalog_products=products,
            buyer=_party("Покупатель"),
            supplier=_party("Поставщик"),
            header=_header(),
            template_bytes=_pricat(),
        )


def test_kit_with_missing_component_is_rejected() -> None:
    products = _products()
    products[2]["components"] = [{"sku": "MISSING", "quantity": 1}]
    with pytest.raises(ValueError, match="не найден компонент"):
        build_vseinstrumenti_pricat(
            _quantities(),
            catalog_products=products,
            buyer=_party("Покупатель"),
            supplier=_party("Поставщик"),
            header=_header(),
            template_bytes=_pricat(),
        )


def test_quantity_template_has_two_columns() -> None:
    sheet = load_workbook(BytesIO(build_vseinstrumenti_quantity_template())).active
    assert sheet["A1"].value == "Название"
    assert sheet["B1"].value == "Количество"


def test_counterparty_gln_round_trip(db_url: str) -> None:
    repo = CrmRepository(db_url)
    repo.init_schema()
    row = repo.create_counterparty(
        {
            "full_name": "Покупатель",
            "inn": "1234567890",
            "kpp": "123456789",
            "gln": "4600000000000",
        }
    )
    assert row.gln == "4600000000000"
    assert repo.counterparty_to_dict(row)["gln"] == "4600000000000"


def test_bundled_pricat_template_is_available() -> None:
    result, stats = build_vseinstrumenti_pricat(
        _quantities(),
        catalog_products=[],
        buyer=_party("Покупатель"),
        supplier=_party("Поставщик"),
        header=_header(),
        actual_date=date(2026, 9, 24),
    )
    book = load_workbook(BytesIO(result), data_only=True)
    sheet = book["Лист 1"]

    assert sheet["C6"].value == "Актуальные остатки 24.09"
    assert stats["rows_written"] == 445
    assert book.sheetnames == ["Лист 1", "Сертификаты и декларации", "Страны", "Единицы измерения"]
    with ZipFile(BytesIO(result)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"dataValidations" in sheet_xml
    assert b"extLst" in sheet_xml

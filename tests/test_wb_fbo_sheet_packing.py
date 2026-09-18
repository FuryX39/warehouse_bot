"""FBO WB new: Excel кабинета, печать ШК, присвоение товара, выгрузка грузомест."""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.catalog_repository import CatalogRepository
from app.crm_repository import CrmRepository
from app.warehouse_users_repository import WarehouseUserRow
from app.wb_fbo_sheet_repository import BOX_ASSIGNED, BOX_PRINTED, WbFboSheetRepository
from app.wb_fbo_sheet_xlsx import fill_boxes_xlsx, parse_boxes_xlsx, parse_goods_xlsx
from app.web.warehouse_tasks_api_auth import TasksApiActor
from app.web.warehouse_wb_fbo_sheet_routes import register_warehouse_wb_fbo_sheet_routes

_GOODS_ROWS = (
    ("4673746970607", 750, "SS958"),
    ("4673746971086", 700, "SS959"),
    ("4673746971437", 700, "SS960"),
    ("4673746971475", 700, "SS961"),
)


def _workbook_bytes(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _goods_xlsx() -> bytes:
    return _workbook_bytes(
        [["Баркод", "Количество, шт", "Артикул поставщика"], *[list(item) for item in _GOODS_ROWS]]
    )


def _boxes_xlsx(*, count: int = 50) -> bytes:
    rows: list[list[object]] = [
        [
            "Баркод товара",
            "Кол-во товаров",
            "ШК короба",
            "Срок годности",
            "ШК короба для печати в стороннем сервисе",
        ]
    ]
    for index in range(count):
        box_id = str(4662171 + index)
        rows.append(["", 0, box_id, "", f"$Ts;0;0;1;box{index};TAS"])
    return _workbook_bytes(rows)


def test_parse_example_wb_tables() -> None:
    goods = parse_goods_xlsx(_goods_xlsx())
    boxes = parse_boxes_xlsx(_boxes_xlsx())
    assert [item.barcode for item in goods] == [
        "4673746970607",
        "4673746971086",
        "4673746971437",
        "4673746971475",
    ]
    assert goods[0].qty == 750
    assert goods[0].sku == "SS958"
    assert len(boxes) == 50
    assert boxes[0].box_id == "4662171"
    assert boxes[0].package_code.startswith("$Ts;")
    assert boxes[0].product_barcode == ""
    assert boxes[0].qty == 0


def test_fill_boxes_xlsx_writes_product_and_qty() -> None:
    original = _boxes_xlsx()
    boxes = parse_boxes_xlsx(original)
    filled = fill_boxes_xlsx(
        original,
        [
            {
                "box_id": boxes[0].box_id,
                "product_barcode": "4673746970607",
                "item_qty": 100,
            }
        ],
    )
    parsed = parse_boxes_xlsx(filled)
    assert parsed[0].product_barcode == "4673746970607"
    assert parsed[0].qty == 100
    assert parsed[0].expiry == ""
    assert parsed[1].product_barcode == ""
    assert parsed[1].qty == 0
    assert parsed[0].package_code == boxes[0].package_code

    orig_wb = load_workbook(BytesIO(original), data_only=True)
    filled_wb = load_workbook(BytesIO(filled), data_only=True)
    try:
        orig_headers = next(orig_wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
        filled_headers = next(filled_wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
    finally:
        orig_wb.close()
        filled_wb.close()
    assert orig_headers == filled_headers
    assert orig_headers[0] == "Баркод товара"
    assert "ШК короба" in orig_headers


def _catalog(db_url: str) -> CatalogRepository:
    crm = CrmRepository(db_url)
    crm.init_schema()
    repo = CatalogRepository(db_url)
    repo.init_schema()
    repo.create_product(
        {
            "name": "SS958",
            "sku": "SS958",
            "code": "00081",
            "is_kit": False,
            "barcodes": [{"barcode": "4673746970607", "label": "", "group": ""}],
            "boxes": [{"barcode": "BOX-SS958-100", "quantity": 100}],
            "components": [],
        }
    )
    repo.create_product(
        {
            "name": "SS959",
            "sku": "SS959",
            "code": "00082",
            "is_kit": False,
            "barcodes": [{"barcode": "4673746971086", "label": "", "group": ""}],
            "boxes": [{"barcode": "BOX-SS959-50", "quantity": 50}],
            "components": [],
        }
    )
    return repo


def _packer() -> WarehouseUserRow:
    return WarehouseUserRow(
        id=7,
        login="packer",
        display_name="packer",
        group_id=None,
        group_name="",
        telegram_nick="",
        is_admin=False,
        is_active=True,
        permissions={},
        created_at_ts=0,
        updated_at_ts=0,
    )


def _client(db_url: str, tmp_path, catalog: CatalogRepository):
    packing = WbFboSheetRepository(db_url, files_data_dir=tmp_path / "wb_fbo_sheet")
    packing.init_schema()
    packer = _packer()
    app = FastAPI()
    register_warehouse_wb_fbo_sheet_routes(
        app,
        packing,
        catalog,
        SimpleNamespace(list_assignee_picker=lambda: [{"id": 7, "display_name": "packer"}]),
        lambda: packer,
        lambda: TasksApiActor(user=packer, via_api_token=False),
        include_manager=True,
        packer_prefixes=("/api/v1/fbo-sheet-packing",),
    )
    return TestClient(app), packing


def _open_pallet(client, job_id: int) -> str:
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-pallets",
        json={"count": 1},
    )
    assert printed.status_code == 200, printed.text
    code = printed.json()["pallets"][0]["pallet_id"]
    opened = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": code},
    )
    assert opened.status_code == 200, opened.text
    assert opened.json()["kind"] == "pallet"
    assert opened.json()["pallet"]["status"] == "open"
    return code


def test_create_print_assign_and_download_boxes_xlsx(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]", "supply_id": "41357389"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job = created.json()["job"]
    assert job["pcs_plan"] == 2850
    assert job["box_total"] == 50
    assert job["box_assigned"] == 0
    job_id = job["id"]

    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 2},
    )
    assert printed.status_code == 200, printed.text
    assert len(printed.json()["boxes"]) == 2
    assert printed.json()["boxes"][0]["status"] == BOX_PRINTED
    assert printed.json()["pdf_base64"]
    first_box = printed.json()["boxes"][0]["box_id"]
    second_box = printed.json()["boxes"][1]["box_id"]
    _open_pallet(client, job_id)

    resolved = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": "BOX-SS958-100"},
    )
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["kind"] == "product"
    assert resolved.json()["suggested_qty"] == 100
    assert resolved.json()["product"]["barcode"] == "4673746970607"

    assigned = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": first_box,
            "product_barcode": "4673746970607",
            "quantity": 100,
        },
    )
    assert assigned.status_code == 200, assigned.text
    assert assigned.json()["box"]["status"] == BOX_ASSIGNED
    assert assigned.json()["qty_warning"] == ""

    leftover = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": second_box,
            "product_barcode": "4673746970607",
            "quantity": 37,
        },
    )
    assert leftover.status_code == 200, leftover.text
    assert "37 товара" in leftover.json()["qty_warning"]
    assert "грузоместе" in leftover.json()["qty_warning"]

    downloaded = client.get(f"/api/warehouse/marketplaces/wb-fbo-new/jobs/{job_id}/boxes.xlsx")
    assert downloaded.status_code == 200
    parsed = parse_boxes_xlsx(downloaded.content)
    by_id = {item.box_id: item for item in parsed}
    assert by_id[first_box].product_barcode == "4673746970607"
    assert by_id[first_box].qty == 100
    assert by_id[first_box].expiry == ""
    assert by_id[second_box].qty == 37
    empty = [item for item in parsed if not item.product_barcode]
    assert len(empty) == 48
    assert packing.get_job(job_id).pcs_assigned == 137


def test_parse_mixed_articles_same_cargo_place() -> None:
    content = _workbook_bytes(
        [
            [
                "Баркод товара",
                "Кол-во товаров",
                "ШК короба",
                "Срок годности",
                "ШК короба для печати в стороннем сервисе",
            ],
            ["4673746970607", 10, "4662171", "", "$Ts;0;0;1;box0;TAS"],
            ["4673746971086", 20, "4662171", "", "$Ts;0;0;1;box0;TAS"],
        ]
    )
    parsed = parse_boxes_xlsx(content)
    assert len(parsed) == 2
    assert parsed[0].box_id == parsed[1].box_id == "4662171"
    assert parsed[0].package_code == parsed[1].package_code
    assert {item.product_barcode: item.qty for item in parsed} == {
        "4673746970607": 10,
        "4673746971086": 20,
    }


def test_parse_rejects_same_box_different_package_code() -> None:
    content = _workbook_bytes(
        [
            [
                "Баркод товара",
                "Кол-во товаров",
                "ШК короба",
                "Срок годности",
                "ШК короба для печати в стороннем сервисе",
            ],
            ["4673746970607", 10, "4662171", "", "$Ts;0;0;1;box0;TAS"],
            ["4673746971086", 20, "4662171", "", "$Ts;0;0;1;box1;TAS"],
        ]
    )
    with pytest.raises(ValueError, match="другим кодом печати"):
        parse_boxes_xlsx(content)


def test_fill_boxes_xlsx_mixed_articles_same_cargo() -> None:
    original = _boxes_xlsx(count=3)
    boxes = parse_boxes_xlsx(original)
    filled = fill_boxes_xlsx(
        original,
        [
            {
                "box_id": boxes[0].box_id,
                "product_barcode": "4673746970607",
                "item_qty": 10,
            },
            {
                "box_id": boxes[0].box_id,
                "product_barcode": "4673746971086",
                "item_qty": 20,
            },
        ],
    )
    parsed = parse_boxes_xlsx(filled)
    same = [item for item in parsed if item.box_id == boxes[0].box_id]
    assert len(same) == 2
    assert {(item.product_barcode, item.qty) for item in same} == {
        ("4673746970607", 10),
        ("4673746971086", 20),
    }
    assert same[0].package_code == same[1].package_code == boxes[0].package_code
    assert len(parsed) == 4


def test_fill_boxes_xlsx_writes_expiry() -> None:
    original = _boxes_xlsx(count=2)
    boxes = parse_boxes_xlsx(original)
    filled = fill_boxes_xlsx(
        original,
        [
            {
                "box_id": boxes[0].box_id,
                "product_barcode": "4673746970607",
                "item_qty": 10,
                "expiry": "17.09.2029",
            }
        ],
    )
    parsed = parse_boxes_xlsx(filled)
    assert parsed[0].expiry == "17.09.2029"
    assert parsed[1].expiry == ""


def _enable_shelf_life(catalog: CatalogRepository, barcode: str, *, years: int = 3) -> None:
    product = catalog.find_product_by_barcode(barcode)
    assert product is not None
    catalog.update_product(
        product.id,
        {
            "name": product.name,
            "sku": product.sku,
            "code": product.code,
            "is_kit": False,
            "has_shelf_life": True,
            "shelf_life_years": years,
            "components": [],
        },
    )


def test_assign_ignores_production_date_without_shelf_life(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=2), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    cargo_id = printed.json()["boxes"][0]["box_id"]
    _open_pallet(client, job_id)
    assigned = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
            "production_date": "2026-09-17",
        },
    )
    assert assigned.status_code == 200, assigned.text
    assert assigned.json()["box"]["items"][0]["expiry"] == ""
    downloaded = client.get(f"/api/warehouse/marketplaces/wb-fbo-new/jobs/{job_id}/boxes.xlsx")
    parsed = parse_boxes_xlsx(downloaded.content)
    row = next(item for item in parsed if item.box_id == cargo_id)
    assert row.expiry == ""
    assert packing.get_job(job_id) is not None


def test_assign_requires_and_writes_expiry_when_shelf_life(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    _enable_shelf_life(catalog, "4673746970607", years=3)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=2), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    cargo_id = printed.json()["boxes"][0]["box_id"]
    _open_pallet(client, job_id)
    missing = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
        },
    )
    assert missing.status_code == 400, missing.text
    assert "дату производства" in missing.json()["detail"]
    resolved = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": "4673746970607"},
    )
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["product"]["has_shelf_life"] is True
    assert resolved.json()["product"]["shelf_life_years"] == 3
    assigned = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
            "production_date": "17.09.2026",
        },
    )
    assert assigned.status_code == 200, assigned.text
    assert assigned.json()["box"]["items"][0]["expiry"] == "17.09.2029"
    downloaded = client.get(f"/api/warehouse/marketplaces/wb-fbo-new/jobs/{job_id}/boxes.xlsx")
    parsed = parse_boxes_xlsx(downloaded.content)
    row = next(item for item in parsed if item.box_id == cargo_id)
    assert row.expiry == "17.09.2029"
    assert packing.get_job(job_id, include_lines=True).pcs_assigned == 10


def test_assign_reuses_sku_expiry_and_rewrites_on_new_date(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    _enable_shelf_life(catalog, "4673746970607", years=3)
    _enable_shelf_life(catalog, "4673746971086", years=3)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=4), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 3},
    )
    assert printed.status_code == 200, printed.text
    first_id, second_id, third_id = [box["box_id"] for box in printed.json()["boxes"]]
    _open_pallet(client, job_id)

    first = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": first_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
            "production_date": "17.09.2026",
        },
    )
    assert first.status_code == 200, first.text
    other = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": first_id,
            "product_barcode": "4673746971086",
            "quantity": 5,
            "production_date": "01.01.2024",
        },
    )
    assert other.status_code == 200, other.text
    reused = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": second_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
        },
    )
    assert reused.status_code == 200, reused.text
    rewritten = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": third_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
            "production_date": "17.09.2025",
        },
    )
    assert rewritten.status_code == 200, rewritten.text

    downloaded = client.get(f"/api/warehouse/marketplaces/wb-fbo-new/jobs/{job_id}/boxes.xlsx")
    parsed = parse_boxes_xlsx(downloaded.content)
    by_id = {}
    for item in parsed:
        by_id.setdefault(item.box_id, []).append(item)
    first_rows = {row.product_barcode: row.expiry for row in by_id[first_id]}
    assert first_rows["4673746970607"] == "17.09.2028"
    assert first_rows["4673746971086"] == "01.01.2027"
    assert by_id[second_id][0].expiry == "17.09.2028"
    assert by_id[third_id][0].expiry == "17.09.2028"
    assert packing.get_job(job_id) is not None


def test_assign_multiple_skus_to_one_cargo_place(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]", "supply_id": "41357389"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=3), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    assert printed.status_code == 200, printed.text
    cargo_id = printed.json()["boxes"][0]["box_id"]
    _open_pallet(client, job_id)

    first = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 100,
        },
    )
    assert first.status_code == 200, first.text
    second = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746971086",
            "quantity": 50,
        },
    )
    assert second.status_code == 200, second.text
    box = second.json()["box"]
    assert box["status"] == BOX_ASSIGNED
    assert len(box["items"]) == 2
    assert {item["product_barcode"]: item["quantity"] for item in box["items"]} == {
        "4673746970607": 100,
        "4673746971086": 50,
    }
    job = packing.get_job(job_id, include_lines=True)
    assert job is not None
    assert job.box_assigned == 1
    assert job.pcs_assigned == 150
    downloaded = client.get(f"/api/warehouse/marketplaces/wb-fbo-new/jobs/{job_id}/boxes.xlsx")
    assert downloaded.status_code == 200
    parsed = parse_boxes_xlsx(downloaded.content)
    rows = [item for item in parsed if item.box_id == cargo_id]
    assert len(rows) == 2
    assert {(item.product_barcode, item.qty) for item in rows} == {
        ("4673746970607", 100),
        ("4673746971086", 50),
    }


def test_create_job_groups_prefilled_mixed_cargo(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    mixed = _workbook_bytes(
        [
            [
                "Баркод товара",
                "Кол-во товаров",
                "ШК короба",
                "Срок годности",
                "ШК короба для печати в стороннем сервисе",
            ],
            ["4673746970607", 10, "4662171", "", "$Ts;0;0;1;box0;TAS"],
            ["4673746971086", 20, "4662171", "", "$Ts;0;0;1;box0;TAS"],
            ["", 0, "4662172", "", "$Ts;0;0;1;box1;TAS"],
        ]
    )
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": (
                "goods.xlsx",
                _goods_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "boxes": (
                "boxes.xlsx",
                mixed,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )
    assert created.status_code == 200, created.text
    job = created.json()["job"]
    assert job["box_total"] == 2
    assert job["box_assigned"] == 1
    assert job["pcs_assigned"] == 30
    stored = packing.get_job(job["id"], include_lines=True)
    assert stored is not None
    assigned = next(box for box in stored.boxes if box.status == BOX_ASSIGNED)
    assert len(assigned.items) == 2


def test_assign_requires_open_pallet(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, _packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=1), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    cargo_id = printed.json()["boxes"][0]["box_id"]
    missing = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
        },
    )
    assert missing.status_code == 400, missing.text
    assert "паллет" in missing.json()["detail"].casefold()


def test_print_open_close_pallet_and_bind_box(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]", "supply_id": "41357389"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=2), xlsx_type),
        },
    )
    job_id = created.json()["job"]["id"]
    boxes = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    ).json()["boxes"]
    cargo_id = boxes[0]["box_id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-pallets",
        json={"count": 2},
    )
    assert printed.status_code == 200, printed.text
    assert len(printed.json()["pallets"]) == 2
    assert printed.json()["pdf_base64"]
    first = printed.json()["pallets"][0]["pallet_id"]
    second = printed.json()["pallets"][1]["pallet_id"]
    assert first.startswith("WBPAL-")
    opened = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": first},
    )
    assert opened.status_code == 200, opened.text
    assert opened.json()["kind"] == "pallet"
    assert opened.json()["job"]["open_pallet"]["pallet_id"] == first
    wrong_close = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/close-pallet",
        json={"barcode": second},
    )
    assert wrong_close.status_code == 400, wrong_close.text
    other = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": second},
    )
    assert other.status_code == 400, other.text
    assigned = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 10,
        },
    )
    assert assigned.status_code == 200, assigned.text
    assert assigned.json()["box"]["pallet_human_id"] == first
    closed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/close-pallet",
        json={"barcode": first},
    )
    assert closed.status_code == 200, closed.text
    assert closed.json()["pallet"]["status"] == "closed"
    assert closed.json()["pallet"]["box_count"] == 1
    job = packing.get_job(job_id, include_lines=True)
    assert job is not None
    assert job.open_pallet is None
    assert job.pallet_closed == 1
    second_open = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/resolve",
        json={"barcode": second},
    )
    assert second_open.status_code == 200, second_open.text
    assert second_open.json()["pallet"]["status"] == "open"


def test_unassign_product_from_cargo_place(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=1), xlsx_type),
        },
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    assert printed.status_code == 200, printed.text
    cargo = printed.json()["boxes"][0]
    cargo_id = cargo["box_id"]
    cargo_pk = cargo["id"]
    _open_pallet(client, job_id)
    first = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746970607",
            "quantity": 100,
        },
    )
    assert first.status_code == 200, first.text
    second = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/assign",
        json={
            "barcode": cargo_id,
            "product_barcode": "4673746971086",
            "quantity": 50,
        },
    )
    assert second.status_code == 200, second.text

    one = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/unassign",
        json={"box_id": cargo_pk, "product_barcode": "4673746970607"},
    )
    assert one.status_code == 200, one.text
    box = one.json()["box"]
    assert box["status"] == BOX_ASSIGNED
    assert {item["product_barcode"]: item["quantity"] for item in box["items"]} == {
        "4673746971086": 50,
    }
    job = packing.get_job(job_id, include_lines=True)
    assert job is not None
    assert job.pcs_assigned == 50
    assert job.box_assigned == 1
    remaining = {item.barcode: item.qty_plan - item.qty_assigned for item in job.products}
    assert remaining["4673746970607"] == 750
    assert remaining["4673746971086"] == 650

    empty = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/unassign",
        json={"barcode": cargo_id},
    )
    assert empty.status_code == 200, empty.text
    box = empty.json()["box"]
    assert box["status"] == BOX_PRINTED
    assert box["items"] == []
    assert box["pallet_human_id"]
    job = packing.get_job(job_id, include_lines=True)
    assert job is not None
    assert job.pcs_assigned == 0
    assert job.box_assigned == 0
    remaining = {item.barcode: item.qty_plan - item.qty_assigned for item in job.products}
    assert remaining["4673746970607"] == 750
    assert remaining["4673746971086"] == 700


def test_unassign_empty_cargo_place_rejected(db_url: str, tmp_path) -> None:
    catalog = _catalog(db_url)
    client, _packing = _client(db_url, tmp_path, catalog)
    xlsx_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    created = client.post(
        "/api/warehouse/marketplaces/wb-fbo-new/jobs",
        data={"packer_user_ids": "[7]"},
        files={
            "goods": ("goods.xlsx", _goods_xlsx(), xlsx_type),
            "boxes": ("boxes.xlsx", _boxes_xlsx(count=1), xlsx_type),
        },
    )
    job_id = created.json()["job"]["id"]
    printed = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/print-boxes",
        json={"count": 1},
    )
    cargo_pk = printed.json()["boxes"][0]["id"]
    missing = client.post(
        f"/api/v1/fbo-sheet-packing/jobs/{job_id}/unassign",
        json={"box_id": cargo_pk},
    )
    assert missing.status_code == 400, missing.text
    assert "нет привязанного" in missing.json()["detail"].casefold()



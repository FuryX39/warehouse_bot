"""Сканер товаров по ШК и Excel-журнал сырых сканов."""

from __future__ import annotations

from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.catalog_repository import CatalogRepository
from app.crm_repository import CrmRepository
from app.marking.cis import GS
from app.scan_log_export import build_scan_log_export, normalize_scan_codes
from app.warehouse_users_repository import WarehouseUserRow
from app.web.warehouse_catalog_routes import register_warehouse_catalog_routes
from app.web.warehouse_tools_routes import register_warehouse_tools_routes


def _repo(db_url: str) -> CatalogRepository:
    crm = CrmRepository(db_url)
    crm.init_schema()
    repo = CatalogRepository(db_url)
    repo.init_schema()
    return repo


def _user() -> WarehouseUserRow:
    return WarehouseUserRow(
        id=1,
        login="admin",
        display_name="Admin",
        group_id=None,
        group_name="",
        telegram_nick="",
        is_admin=True,
        is_active=True,
        permissions={},
        created_at_ts=0,
        updated_at_ts=0,
    )


def test_find_product_by_barcode_exact_and_casefold(db_url: str) -> None:
    repo = _repo(db_url)
    product = repo.create_product(
        {
            "name": "Болт",
            "sku": "BOLT-1",
            "code": "00001",
            "is_kit": False,
            "barcodes": [{"barcode": "ABC-12345", "label": "", "group": ""}],
            "components": [],
        }
    )
    found = repo.find_product_by_barcode("ABC-12345")
    assert found is not None
    assert found.id == product.id
    assert found.sku == "BOLT-1"
    assert repo.find_product_by_barcode("abc-12345") is not None
    assert repo.find_product_by_barcode("MISSING") is None
    assert repo.find_product_by_barcode("   ") is None


def test_normalize_scan_codes_keeps_duplicates_and_gs() -> None:
    codes = normalize_scan_codes({"codes": ["AAA", "AAA", f"01{GS}21"]})
    assert codes == ["AAA", "AAA", f"01{GS}21"]


def test_build_scan_log_export_writes_raw_rows() -> None:
    raw = build_scan_log_export(["AAA", f"01{GS}21", "AAA"])
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    assert ws.title == "Сканы"
    assert ws.cell(1, 1).value == "№"
    assert ws.cell(1, 2).value == "Скан"
    assert ws.cell(2, 1).value == 1
    assert ws.cell(2, 2).value == "AAA"
    assert ws.cell(3, 2).value == "01<GS>21"
    assert ws.cell(4, 2).value == "AAA"


def test_catalog_by_barcode_http(db_url: str) -> None:
    repo = _repo(db_url)
    repo.create_product(
        {
            "name": "Гайка",
            "sku": "NUT-1",
            "code": "00002",
            "is_kit": False,
            "barcodes": [{"barcode": "NUT-999", "label": "осн", "group": ""}],
            "components": [],
        }
    )
    app = FastAPI()

    def require_warehouse_user() -> WarehouseUserRow:
        return _user()

    register_warehouse_catalog_routes(app, repo, require_warehouse_user)
    client = TestClient(app)
    empty = client.get("/api/warehouse/catalog/products/by-barcode")
    assert empty.status_code == 400
    miss = client.get("/api/warehouse/catalog/products/by-barcode", params={"barcode": "NOPE"})
    assert miss.status_code == 200
    assert miss.json()["product"] is None
    hit = client.get("/api/warehouse/catalog/products/by-barcode", params={"barcode": "nut-999"})
    assert hit.status_code == 200, hit.text
    product = hit.json()["product"]
    assert product["sku"] == "NUT-1"
    assert product["name"] == "Гайка"
    assert any(b["barcode"] == "NUT-999" for b in product["barcodes"])


def test_tools_scanner_export_http() -> None:
    app = FastAPI()

    def require_warehouse_user() -> WarehouseUserRow:
        return _user()

    register_warehouse_tools_routes(app, catalog_repo=None, require_warehouse_user=require_warehouse_user)
    client = TestClient(app)
    empty = client.post("/api/warehouse/tools/scanner/export", json={"codes": []})
    assert empty.status_code == 400
    exported = client.post(
        "/api/warehouse/tools/scanner/export",
        json={"codes": ["FILE-1", "FILE-1", f"DM{GS}TAIL"]},
    )
    assert exported.status_code == 200, exported.text
    wb = load_workbook(BytesIO(exported.content))
    ws = wb.active
    assert [ws.cell(2, 2).value, ws.cell(3, 2).value, ws.cell(4, 2).value] == [
        "FILE-1",
        "FILE-1",
        "DM<GS>TAIL",
    ]

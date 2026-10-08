"""Отчёт по продажам: период, маркетплейс, сумма из «Цена с НДС»."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.catalog_repository import CatalogProduct, CatalogRepository
from app.crm_repository import CrmRepository
from app.sales_report import build_sales_report, parse_report_dates
from app.storage_warehouse_repository import StorageWarehouseRepository
from app.warehouse_orders_repository import (
    ORDER_CANCELLED,
    SOURCE_WB,
    SOURCE_YM,
    WarehouseOrder,
    WarehouseOrdersRepository,
)

_TZ = ZoneInfo("Europe/Moscow")
_DAY = datetime(2026, 10, 5, 12, 0, tzinfo=_TZ)
_TS = int(_DAY.timestamp())


def _stack(db_url: str):
    storage = StorageWarehouseRepository(db_url)
    storage.init_schema()
    catalog = CatalogRepository(db_url)
    catalog.init_schema()
    crm = CrmRepository(db_url)
    crm.init_schema()
    orders = WarehouseOrdersRepository(db_url)
    orders.init_schema()
    return storage, catalog, orders, int(storage.get_default_warehouse_id())


def _add_product(catalog: CatalogRepository, sku: str, name: str) -> None:
    with Session(catalog.engine) as session:
        session.add(
            CatalogProduct(name=name, sku=sku, code=sku[:8], created_at_ts=1, updated_at_ts=1)
        )
        session.commit()


def _set_created(orders: WarehouseOrdersRepository, order_id: int, ts: int, status: str | None = None) -> None:
    with Session(orders.engine) as session:
        row = session.get(WarehouseOrder, int(order_id))
        assert row is not None
        row.created_at_ts = int(ts)
        if status is not None:
            row.status = status
        session.commit()


def test_parse_report_dates_rejects_inverted_range() -> None:
    with pytest.raises(ValueError, match="раньше"):
        parse_report_dates("2026-10-08", "2026-10-01")


def test_yandex_sales_report_sums_unit_price(db_url: str) -> None:
    _storage, catalog, orders, wh_id = _stack(db_url)
    _add_product(catalog, "SS100", "NanoGlass")
    _add_product(catalog, "SS200", "NanoSet")
    first = orders.upsert_from_posting(
        source=SOURCE_YM,
        posting_id="ym-1",
        warehouse_id=wh_id,
        lines=[{"sku": "SS100", "quantity": 2, "name": ""}],
        now_ts=_TS,
    )
    orders.apply_empty_line_money(first.id, [{"sku": "SS100", "unit_price": "374"}])
    _set_created(orders, first.id, _TS)
    second = orders.upsert_from_posting(
        source=SOURCE_YM,
        posting_id="ym-2",
        warehouse_id=wh_id,
        lines=[
            {"sku": "SS100", "quantity": 1, "name": ""},
            {"sku": "SS200", "quantity": 4, "name": ""},
        ],
        now_ts=_TS,
    )
    orders.apply_empty_line_money(
        second.id,
        [{"sku": "SS100", "unit_price": "374"}, {"sku": "SS200", "unit_price": "50"}],
    )
    _set_created(orders, second.id, _TS)
    cancelled = orders.upsert_from_posting(
        source=SOURCE_YM,
        posting_id="ym-cancel",
        warehouse_id=wh_id,
        lines=[{"sku": "SS100", "quantity": 99, "name": ""}],
        now_ts=_TS,
    )
    orders.apply_empty_line_money(cancelled.id, [{"sku": "SS100", "unit_price": "374"}])
    _set_created(orders, cancelled.id, _TS, status=ORDER_CANCELLED)
    outside = orders.upsert_from_posting(
        source=SOURCE_YM,
        posting_id="ym-old",
        warehouse_id=wh_id,
        lines=[{"sku": "SS200", "quantity": 10, "name": ""}],
        now_ts=_TS,
    )
    orders.apply_empty_line_money(outside.id, [{"sku": "SS200", "unit_price": "50"}])
    _set_created(orders, outside.id, int(datetime(2026, 9, 1, 12, tzinfo=_TZ).timestamp()))

    result = build_sales_report(
        orders,
        catalog,
        marketplace_id="yandex_market",
        date_from="2026-10-01",
        date_to="2026-10-08",
    )
    assert [(row.sku, row.name, row.quantity, row.amount) for row in result.rows] == [
        ("SS100", "NanoGlass", 3, Decimal("1122.00")),
        ("SS200", "NanoSet", 4, Decimal("200.00")),
    ]
    assert result.total_quantity == 7
    assert result.total_amount == Decimal("1322.00")
    assert result.missing_amount_count == 0
    book = load_workbook(BytesIO(result.workbook_bytes))
    sheet = book.active
    assert [cell.value for cell in sheet[4]] == ["Артикул", "Название", "Количество", "Сумма"]
    assert sheet[5][3].value == 1122.0
    assert sheet[7][0].value == "Итого"
    assert sheet[7][3].value == 1322.0


def test_wb_sales_report_leaves_sum_empty(db_url: str) -> None:
    _storage, catalog, orders, wh_id = _stack(db_url)
    _add_product(catalog, "SS100", "NanoGlass")
    row = orders.upsert_from_posting(
        source=SOURCE_WB,
        posting_id="wb-1",
        warehouse_id=wh_id,
        lines=[{"sku": "SS100", "quantity": 5, "name": ""}],
        now_ts=_TS,
    )
    _set_created(orders, row.id, _TS)
    result = build_sales_report(
        orders,
        catalog,
        marketplace_id="wildberries",
        date_from="2026-10-01",
        date_to="2026-10-08",
    )
    assert len(result.rows) == 1
    assert result.rows[0].sku == "SS100"
    assert result.rows[0].name == "NanoGlass"
    assert result.rows[0].quantity == 5
    assert result.rows[0].amount is None
    assert result.total_amount is None
    assert result.missing_amount_count == 1
    book = load_workbook(BytesIO(result.workbook_bytes))
    sheet = book.active
    assert sheet[5][2].value == 5
    assert sheet[5][3].value is None
    assert sheet[6][0].value == "Итого"
    assert sheet[6][3].value is None

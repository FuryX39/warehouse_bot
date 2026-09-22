"""Заказ ВсеИнструменты: разбор Excel, пик и лист несовпадений ШК."""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.catalog_repository import CatalogRepository
from app.crm_repository import CrmRepository
from app.other_marketplace_repository import OtherMarketplaceRepository
from app.other_marketplace_service import (
    build_other_marketplace_marking_xlsx,
    create_vseinstrumenti_job,
    pick_other_marketplace_line,
)
from app.vseinstrumenti_order import parse_vseinstrumenti_order


def _xlsx(*, header_row: int = 18) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 'Подтверждение заказа №4394602 от 22.09.2026 ООО "ВсеИнструменты.ру"'
    ws["A4"] = "Ожидаемое время доставки по заказу:"
    ws["D4"] = "10.10.2026 00:00"
    ws["A14"] = "Поставщик:"
    ws["C14"] = 'ООО "ШАЙН СИСТЕМС"'
    ws.cell(header_row - 2, 1, "Артикул")
    ws.cell(10, 3, "Наименование")
    headers = ["№", "Штрихкод", "Наименование", "Код сети", "Артикул", "Заказано", "Подтверждено"]
    for col, title in enumerate(headers, start=1):
        ws.cell(header_row, col, title)
    rows = [
        (1, "4673746970683", "Паста из файла", "SS585", "SS585", 5, 5),
        (2, "9990000000002", "Другой штрихкод", "SS100", "SS100", 2, 2),
    ]
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            ws.cell(header_row + 1 + offset, col, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _repos(db_url: str):
    CrmRepository(db_url).init_schema()
    catalog = CatalogRepository(db_url)
    catalog.init_schema()
    repo = OtherMarketplaceRepository(db_url)
    repo.init_schema()
    return catalog, repo


def test_parser_reads_order_barcode_and_confirmed_qty() -> None:
    order = parse_vseinstrumenti_order(_xlsx())
    assert order.order_number == "4394602"
    assert order.delivery_date == "2026-10-10"
    assert order.supplier.startswith("ООО")
    assert [(line.sku, line.excel_barcode, line.quantity) for line in order.lines] == [
        ("SS585", "4673746970683", 5),
        ("SS100", "9990000000002", 2),
    ]


def test_product_list_follows_header_above_or_below_row_18() -> None:
    expected = [
        ("SS585", "4673746970683", 5),
        ("SS100", "9990000000002", 2),
    ]
    for header_row in (12, 18, 30):
        order = parse_vseinstrumenti_order(_xlsx(header_row=header_row))
        assert order.order_number == "4394602"
        assert order.delivery_date == "2026-10-10"
        assert [(line.sku, line.excel_barcode, line.quantity) for line in order.lines] == expected


def test_real_example_barcode_is_column_b() -> None:
    folder = Path(r"C:\3d\warehouse_bot\tables_examples") / "ВИ"
    files = [path for path in folder.glob("*.xlsx") if not path.name.startswith("~$")]
    if not files:
        return
    order = parse_vseinstrumenti_order(files[0].read_bytes())
    assert order.order_number
    assert order.lines
    assert order.lines[0].excel_barcode.isdigit()


def test_catalog_barcode_mismatch_goes_to_manager_sheet(db_url: str) -> None:
    catalog, repo = _repos(db_url)
    catalog.create_product(
        {
            "name": "EasyFinish из карточки",
            "sku": "SS585",
            "code": "585",
            "barcodes": [{"barcode": "1111222233334", "label": ""}],
        }
    )
    catalog.create_product(
        {
            "name": "NanoGlass из карточки",
            "sku": "SS100",
            "code": "100",
            "barcodes": [{"barcode": "5555666677778", "label": ""}],
        }
    )
    job, warnings = create_vseinstrumenti_job(
        catalog=catalog,
        repo=repo,
        content=_xlsx(),
        filename="order.xlsx",
        transfer_number="ПР-15",
        packer_user_ids=[1],
        created_by_user_id=1,
    )
    assert warnings == []
    assert job.lines[0].product_name == "EasyFinish из карточки"
    assert job.lines[0].excel_barcode == "4673746970683"

    same = pick_other_marketplace_line(
        catalog=catalog,
        repo=repo,
        job_id=job.id,
        raw="4673746970683",
    )
    assert same["mismatch"] is False
    assert same["barcode"] == "4673746970683"
    assert same["barcode_copies"] == 5

    other = pick_other_marketplace_line(
        catalog=catalog,
        repo=repo,
        job_id=job.id,
        raw="5555666677778",
    )
    assert other["mismatch"] is True
    assert other["barcode"] == "9990000000002"
    assert other["barcode_copies"] == 2

    saved = repo.get_job(job.id)
    assert saved is not None
    content = build_other_marketplace_marking_xlsx(saved)
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["Все строки", "С КИЗ", "Несовпадения ШК"]
    sheet = book["Несовпадения ШК"]
    assert sheet["A2"].value == "4394602"
    assert sheet["B2"].value == "SS100"
    assert sheet["D2"].value == "9990000000002"
    assert sheet["E2"].value == "5555666677778"
    assert sheet.max_row == 2
